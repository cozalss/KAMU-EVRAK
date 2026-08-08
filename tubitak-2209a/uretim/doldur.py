"""Resmi TUBITAK 2209-A Arastirma Onerisi Formu'nu icerik.json ile doldurur.

Resmi formun kendi bicimini (Arial 9, tablo yapisi, kenar bosluklari) KORUR;
sadece bos hucreleri ve yer tutuculari gercek metinle degistirir.

Form haritasi (formcheck.py ile olculdu):
  T0 4x1  Genel bilgiler (ad, baslik, danisman, kurum)
  T1 2x1  r0=Ozet  r1=Anahtar Kelimeler
  T2 1x1  1.1 Konunun Onemi ve Bilimsel Nitelik
  T3 1x1  1.2 Amac ve Hedefler
  T4 1x1  2. Yontem
  T5 6x4  3.1 Calisma Takvimi (r0 baslik)
  T6 3x2  3.2 Risk Yonetimi (r0 baslik)
  T7 3x2  3.3 Arastirma Olanaklari (r0 baslik)
  T8 4x2  4. Yaygin Etki (r0 baslik, r1..r3 kategoriler)
  T9 1x1  5. Belirtmek Istediginiz Diger Konular
  EK-1 KAYNAKLAR: son paragraftan sonra eklenir
"""
import sys, os, json, copy
sys.stdout.reconfigure(encoding="utf-8")

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from bicim import zengin_yaz
from akis import zengin_akis, faktoriyel_akis, gantt_akis

# Resmi form sablonu once betikle AYNI klasorde, sonra bir ust klasorde aranir.
# (8 Agu: sablon uretim/ icine tasinmisti ama yol bir ust klasoru gosteriyordu;
# doldur.py bu yuzden hic kosmuyordu — teslim belgesi kaynaktan yeniden
# uretilemez durumdaydi. Yol artik iki konumu da dener ve bulamazsa ACIK soyler.)
_BURASI = os.path.dirname(os.path.abspath(__file__))
_ADAYLAR = (os.path.join(_BURASI, "2209A_resmi_form.docx"),
            os.path.join(os.path.dirname(_BURASI), "2209A_resmi_form.docx"))
FORM = next((y for y in _ADAYLAR if os.path.exists(y)), None)
if FORM is None:
    raise SystemExit("Resmî form şablonu bulunamadı. Aranan konumlar:\n  " +
                     "\n  ".join(_ADAYLAR))

ARIAL = "Arial"
BOYUT = Pt(9)


# --------------------------------------------------------------- yardimcilar
def _bicimle(run, bold=False, italic=False, size=BOYUT):
    run.font.name = ARIAL
    run.font.size = size
    run.font.bold = bold
    run.font.italic = italic
    # eastAsia/cs ayarlari: Turkce karakterlerin tum gorunumlerde Arial kalmasi icin
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
            + attr.split(":")[1],
            ARIAL,
        )
    return run


def hucre_yaz(cell, bloklar, justify=True):
    """Hucreyi temizler ve bloklari yazar.

    bloklar: str  ya da  [(metin, {"bold":..,"italic":..,"bullet":..}), ...]
    """
    if isinstance(bloklar, str):
        bloklar = [b for b in bloklar.split("\n")]
    cell._element.clear_content()
    yazildi = False
    for blok in bloklar:
        if isinstance(blok, tuple):
            metin, opt = blok
        else:
            metin, opt = blok, {}
        metin = (metin or "").strip()
        p = cell.add_paragraph()
        pf = p.paragraph_format
        # 20 sayfa siniri dar: paragraf arasi 3pt yerine 2pt. Icerikten odun
        # vermeden yaklasik yarim sayfa kazandirir, okunurlugu bozmaz.
        pf.space_after = Pt(2)
        pf.space_before = Pt(0)
        if justify and not opt.get("bold") and len(metin) > 120:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if not metin:
            _bicimle(p.add_run(""))
            continue
        if opt.get("bullet"):
            metin = "•  " + metin
            pf.left_indent = Pt(10)
        _bicimle(p.add_run(metin), bold=opt.get("bold", False),
                 italic=opt.get("italic", False), size=opt.get("size", BOYUT))
        yazildi = True
    if not yazildi and not cell.paragraphs:
        _bicimle(cell.add_paragraph().add_run(""))


def hucre_ekle_deger(cell, deger):
    """Var olan etiketin (orn. 'Danismanin Adi Soyadi:') yanina degeri ekler."""
    p = cell.paragraphs[-1] if cell.paragraphs else cell.add_paragraph()
    _bicimle(p.add_run("  " + str(deger)), bold=False)


def satir_kopyala(tablo):
    """Tablonun son satirini kopyalayip sonuna ekler ve yeni satiri dondurur."""
    son = tablo.rows[-1]._tr
    yeni = copy.deepcopy(son)
    son.addnext(yeni)
    return tablo.rows[-1]


def satir_garanti(tablo, gerekli_veri_satiri, baslik_satiri=1):
    """Tabloda yeterli veri satiri olmasini saglar (baslik satirlari haric)."""
    mevcut = len(tablo.rows) - baslik_satiri
    for _ in range(max(0, gerekli_veri_satiri - mevcut)):
        satir_kopyala(tablo)
    return tablo


def paragraf_ekle_sonuna(doc, metin, bold=False, size=BOYUT):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    _bicimle(p.add_run(metin), bold=bold, size=size)
    return p


def paragraf_metnini_degistir(p, yeni, bold=False):
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    _bicimle(p.add_run(yeni), bold=bold)


# --------------------------------------------------------------------- ana
def doldur(icerik, cikti):
    doc = Document(FORM)
    T = doc.tables

    # --- baslik satirlari: yil / donem
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith(".....Yılı"):
            paragraf_metnini_degistir(p, icerik.get("yil", "2026") + " Yılı")
        elif t.startswith("… Dönem"):
            # 2209-A'da DONEM numaralandirmasi kaldirildi (cagri artik "<yil> Yili
            # Cagrisi" diye aniliyor). donem bos ise satir tamamen bosaltilir.
            d = str(icerik.get("donem", "")).strip()
            paragraf_metnini_degistir(p, f"{d} Dönem Başvurusu" if d else "")

    # --- T0 genel bilgiler
    genel = icerik.get("genel_bilgiler", {})
    for satir, anahtar in zip(T[0].rows, ["ad_soyad", "baslik", "danisman", "kurum"]):
        deger = genel.get(anahtar)
        if deger:
            hucre_ekle_deger(satir.cells[0], deger)

    # --- T1 ozet + anahtar kelimeler
    ozet = icerik.get("ozet", "").strip()
    if ozet:
        bloklar = [("Özet", {"bold": True})]
        bloklar += [par for par in ozet.split("\n") if par.strip()]
        hucre_yaz(T[1].rows[0].cells[0], bloklar)
    kelimeler = icerik.get("anahtar_kelimeler", [])
    if kelimeler:
        c = T[1].rows[1].cells[0]
        c._element.clear_content()
        p = c.add_paragraph()
        _bicimle(p.add_run("Anahtar Kelimeler: "), bold=True)
        _bicimle(p.add_run(", ".join(kelimeler)))

    # --- T2 / T3 / T4 serbest metin bolumleri
    # Bu kutular birer GIRDI ALANIDIR; sayfalarca suren metin icin kap degildir.
    # Icerik en sonda govde akisina dokulur (bkz. akisa_cevir listesi).
    akisa_cevir = []
    for idx, anahtar in ((2, "bolum_1_1"), (3, "bolum_1_2"), (4, "bolum_2")):
        metin = icerik.get(anahtar, "").strip()
        if metin:
            akisa_cevir.append((T[idx], metin))

    # --- 3. PROJE YONETIMI giris metni (rubrikte %20; hakem "bu blok uc tablodan
    # ibaret, izleme/koordinasyon yordami hicbir yerde yok" dedi). Baslikla ilk alt
    # baslik arasina akis olarak yerlestirilir.
    yonetim = icerik.get("bolum_3_yonetim_girisi", "").strip()
    if yonetim:
        for p in doc.paragraphs:
            if p.text.strip().startswith("PROJE YÖNETİMİ"):
                # basligin ARDINDAN gelen ilk ogeden once yaz
                sonraki = p._p.getnext()
                if sonraki is not None:
                    zengin_akis(doc, sonraki, yonetim)
                break

    # --- T5 calisma takvimi
    takvim = icerik.get("calisma_takvimi", [])
    if takvim:
        satir_garanti(T[5], len(takvim), baslik_satiri=1)
        for i, ip in enumerate(takvim, start=1):
            hucreler = T[5].rows[i].cells
            hucre_yaz(hucreler[0], str(ip.get("tarih_araligi", "")), justify=False)
            fa = ip.get("faaliyetler", "")
            ad = ip.get("ad")
            bloklar = ([(ad, {"bold": True})] if ad else []) + \
                      [x for x in str(fa).split("\n") if x.strip()]
            # Dar sutunda iki yana yaslama kelime araliklarini acip "nehir" olusturuyor;
            # takvim hucrelerinde sola yaslama okunakli.
            hucre_yaz(hucreler[1], bloklar, justify=False)
            hucre_yaz(hucreler[2], str(ip.get("kimler", "")), justify=False)
            olcut = str(ip.get("basari_olcutu", ""))
            katki = ip.get("katki")
            if katki is not None:
                olcut = f"{olcut}\nKatkı: %{katki}"
            hucre_yaz(hucreler[3], olcut, justify=False)
        # fazla bos satirlari sil
        for r in list(T[5].rows[len(takvim) + 1:]):
            r._tr.getparent().remove(r._tr)

    # --- T6 riskler
    riskler = icerik.get("riskler", [])
    if riskler:
        satir_garanti(T[6], len(riskler), baslik_satiri=1)
        for i, r in enumerate(riskler, start=1):
            hucre_yaz(T[6].rows[i].cells[0], str(r.get("risk", "")))
            hucre_yaz(T[6].rows[i].cells[1], str(r.get("tedbir", "")))
        for row in list(T[6].rows[len(riskler) + 1:]):
            row._tr.getparent().remove(row._tr)

    # --- T7 arastirma olanaklari
    olanaklar = icerik.get("olanaklar", [])
    if olanaklar:
        satir_garanti(T[7], len(olanaklar), baslik_satiri=1)
        for i, o in enumerate(olanaklar, start=1):
            hucre_yaz(T[7].rows[i].cells[0], str(o.get("altyapi", "")))
            hucre_yaz(T[7].rows[i].cells[1], str(o.get("kullanim_amaci", "")))
        for row in list(T[7].rows[len(olanaklar) + 1:]):
            row._tr.getparent().remove(row._tr)

    # --- T8 yaygin etki (3 kategori, sag sutun)
    etki = icerik.get("yaygin_etki", {})
    for i, anahtar in enumerate(
            ["bilimsel_akademik", "ekonomik_ticari_sosyal", "yeni_projeler"], start=1):
        metin = etki.get(anahtar, "").strip()
        if metin:
            hucre_yaz(T[8].rows[i].cells[1],
                      [x for x in metin.split("\n") if x.strip()])

    # --- T9 diger konular + SEKILLER
    # Form bu bolum icin acikca "grafik, tablo vb. eklenebilir" diyor; iki sekil
    # buraya konur. PNG degil YERLI TABLO: her yakinlastirmada net, s/b basilir,
    # Word'de duzenlenebilir, dis bagimlilik yok.
    diger = icerik.get("bolum_5_diger", "").strip()
    if diger:
        akisa_cevir.append((T[9], diger))

    # NOT: Ekler oncesine sayfa sonu KONULMUYOR.
    # Govde artik ~17 sayfa; 20 sinirinin cok altinda oldugu icin ekleri ayri sayfaya
    # itmenin bir koruma degeri yok, buna karsilik tam bir sayfa harciyor. Kullanicinin
    # istegi "ekler dahil 19 sayfa" oldugundan bu sayfa geri kazanildi.

    # --- EK-1 KAYNAKLAR
    kaynaklar = icerik.get("kaynaklar", [])
    if kaynaklar:
        for k in kaynaklar:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Pt(18)
            p.paragraph_format.first_line_indent = Pt(-18)
            _bicimle(p.add_run(k))

    # --- EK-2 BUTCE (form zorunlu kilmiyor; degerlendirmeye katki icin)
    butce = icerik.get("butce", [])
    if butce:
        paragraf_ekle_sonuna(doc, "", bold=False)
        paragraf_ekle_sonuna(doc, "EK-2:  TAHMİNİ BÜTÇE", bold=True)
        tb = doc.add_table(rows=1, cols=3)
        tb.style = "Table Grid"
        basliklar = ["Gider Kalemi", "Açıklama ve Gerekçe", "Tutar (TL)"]
        for c, b in zip(tb.rows[0].cells, basliklar):
            hucre_yaz(c, [(b, {"bold": True})], justify=False)
        toplam = 0
        for kalem in butce:
            row = tb.add_row()
            hucre_yaz(row.cells[0], str(kalem.get("kalem", "")), justify=False)
            hucre_yaz(row.cells[1], str(kalem.get("aciklama", "")))
            tutar = kalem.get("tutar_tl", 0) or 0
            toplam += tutar
            hucre_yaz(row.cells[2], f"{tutar:,.0f}".replace(",", "."), justify=False)
        row = tb.add_row()
        hucre_yaz(row.cells[0], [("TOPLAM", {"bold": True})], justify=False)
        hucre_yaz(row.cells[1], "")
        hucre_yaz(row.cells[2],
                  [(f"{toplam:,.0f}".replace(",", "."), {"bold": True})], justify=False)

    # --- GÖRSEL GEÇİŞ, 1. aşama: form biçimi (kutular hâlâ yerindeyken)
    # NOT: tablo NESNELERİ geçiliyor, indeks değil — kutular kaldırılınca
    # doc.tables indeksleri kayar ve indeksle çalışmak yanlış tabloyu biçimler.
    from tasarim import uygula_form, uygula_son
    olcum = uygula_form(doc, [
        (T[5], True, True, True),    # çalışma takvimi (sayfa aşar → başlık yinelensin)
        (T[6], True, True, False),   # riskler
        (T[7], True, True, False),   # araştırma olanakları
        (T[8], True, False, False),  # yaygın etki
    ])

    # --- 2. aşama: içerik kutularını gövde akışına dök
    for tablo, metin in akisa_cevir:
        hedef = tablo._tbl
        zengin_akis(doc, hedef, metin)
        if tablo is T[9]:
            faktoriyel_akis(doc, hedef)
            gantt_akis(doc, hedef, icerik.get("calisma_takvimi", []))
        hedef.getparent().remove(hedef)

    # --- 3. aşama: son rötuş (kaynakça, dul/yetim satır)
    olcum.update(uygula_son(doc))
    print("  GÖRSEL:", olcum)

    doc.save(cikti)
    return cikti


if __name__ == "__main__":
    icerik_yolu = sys.argv[1]
    cikti_yolu = sys.argv[2]
    with open(icerik_yolu, encoding="utf-8") as f:
        icerik = json.load(f)
    yol = doldur(icerik, cikti_yolu)
    print("YAZILDI:", yol)
