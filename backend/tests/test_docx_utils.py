"""docx tablo metni çıkarımı — hücre KİMLİĞİ (id) bazlı teksilleştirme regresyonu.

ed9a8666 (bug avı 2. dalga, HIGH) bulgusu: eski kod `dict.fromkeys` ile hücre metin
DEĞERİne göre dedup ediyordu → farklı sütunların tesadüfen aynı metni ('Evet', '0 TL',
'-') sessizce düşüyordu (veri kaybı). Düzeltme: birleşik olmayan hücreler ayrı `_tc`
XML elemanı taşıdığından `id(c._tc)` bazlı dedup — yalnız gerçek merged hücreler tekilleşir.
"""

from __future__ import annotations

import io

import docx

from multimodal.docx_utils import extract_docx_text


def _docx_bytes(document) -> bytes:
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


class TestDocxTabloDedup:
    def test_farkli_hucreler_ayni_metinle_dusurulmez(self):
        # Arrange — 3 ayrı hücre, 0 ve 2 tesadüfen aynı metin ("Evet")
        document = docx.Document()
        table = document.add_table(rows=1, cols=3)
        table.rows[0].cells[0].text = "Evet"
        table.rows[0].cells[1].text = "Hayır"
        table.rows[0].cells[2].text = "Evet"

        # Act
        metin = extract_docx_text(_docx_bytes(document))

        # Assert — üç hücre de korunur (pre-fix dict.fromkeys 2. "Evet"i düşürürdü)
        assert "Evet | Hayır | Evet" in metin

    def test_gercek_birlesik_hucre_tek_kez_gelir(self):
        # Arrange — 0 ve 1 GERÇEKTEN birleştirildi → aynı _tc → tek metin
        document = docx.Document()
        table = document.add_table(rows=1, cols=3)
        table.rows[0].cells[0].text = "Birleşik"
        table.rows[0].cells[2].text = "Ayrı"
        table.rows[0].cells[0].merge(table.rows[0].cells[1])

        # Act
        metin = extract_docx_text(_docx_bytes(document))

        # Assert — birleşik hücre bir kez; ayrı hücre korunur
        assert "Birleşik | Ayrı" in metin
        assert "Birleşik | Birleşik" not in metin
