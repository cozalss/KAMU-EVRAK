"""Teslim belgesinde doldurulmayi bekleyen tum yer tutuculari bulur."""
import sys, re
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document

doc = Document(sys.argv[1])

# yer tutucu kaliplari: uc nokta dizisi, nokta dizisi, kose parantezli not
DESENLER = [
    (r"…{3,}", "üç nokta dizisi"),
    (r"\.{4,}", "nokta dizisi"),
    (r"\[DOLDUR[^\]]*\]", "DOLDUR etiketi"),
    (r"\[DOĞRULA[^\]]*\]", "DOĞRULA etiketi"),
    (r"20\.\.", "yıl yer tutucu"),
]

bulundu = []


def tara(kaynak, nereden):
    for desen, ad in DESENLER:
        for m in re.finditer(desen, kaynak):
            a, b = max(0, m.start() - 130), min(len(kaynak), m.end() + 130)
            bulundu.append((nereden, ad, kaynak[a:b].replace("\n", " ")))


TABLO_ADI = {}
for i, t in enumerate(doc.tables):
    ilk = " ".join(c.text for c in t.rows[0].cells)[:45].replace("\n", " ")
    TABLO_ADI[i] = f"T{i} ({ilk})"

for i, t in enumerate(doc.tables):
    for ri, r in enumerate(t.rows):
        for ci, c in enumerate(r.cells):
            tara(c.text, f"{TABLO_ADI[i]} r{ri}c{ci}")

for pi, p in enumerate(doc.paragraphs):
    tara(p.text, f"paragraf {pi}")

print("=" * 78)
print(f"DOLDURULACAK ALANLAR — {len(bulundu)} adet")
print("=" * 78)
for nereden, ad, baglam in bulundu:
    print(f"\n[{ad}] {nereden}")
    print(f"   …{baglam.strip()}…")
