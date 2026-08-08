"""B. 2209A_REHBERLI_SABLON.docx uretir.

Her form bolumu icin dort katman:
  [TUBITAK]            formun kendi resmi aciklamasi (birebir alinti)
  [PANELIST NE ARAR]   rubrik ipucu + puan agirligi + puan kesme riski
  [ORNEK]              icerik.json'daki gercek doldurulmus metinden alinti
  [SIZ YAZIN]          bos alan + karakter hedefi

Kullanim: python sablon_yap.py <icerik.json> <cikti.docx>
"""
import sys, os, json, re
sys.stdout.reconfigure(encoding="utf-8")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from ekler import sayfa_kur, baslik, gövde, kutu, bos_alan, tablo_yap, _font


def alinti(metin, azami=1500):
    """Ornek metinden temsili bir parca alir (paragraf sinirinda keser)."""
    metin = (metin or "").strip()
    if not metin:
        return ["(Bu bölüm için örnek metin üretilmedi.)"]
    if len(metin) <= azami:
        return [p for p in metin.split("\n") if p.strip()]
    kesik = metin[:azami]
    son = max(kesik.rfind("\n"), kesik.rfind(". "))
    if son > azami * 0.5:
        kesik = kesik[:son + 1]
    parcalar = [p for p in kesik.split("\n") if p.strip()]
    parcalar.append("[...] — Tam metnin devamı TESLİM dosyasındadır.")
    return parcalar


BOLUM_REHBERI = [
    dict(
        no="A", ad="GENEL BİLGİLER", hedef=0, tur="tablo",
        tubitak=["Form dört alan ister: Başvuru Sahibinin Adı Soyadı, Araştırma Önerisinin Başlığı, "
                 "Danışmanın Adı Soyadı, Araştırmanın Yürütüleceği Kurum/Kuruluş."],
        panelist=["Ağırlık: doğrudan puanı yok — ama başlık panelistin gördüğü İLK şeydir.",
                  "Madde 4.1.7: proje başlığı araştırmanın amacını ve içeriğini yansıtan, açık ve "
                  "anlaşılabilir bilimsel ifadelerle belirlenmelidir.",
                  "Madde 4.1.8: başlıkta hiçbir kurum, firma, ticari ürün veya marka geçemez.",
                  "İyi başlık kalıbı: <yöntem/yaklaşım> ile <problem alanı>: <ne ölçüleceği>. "
                  "Kısaltma, slogan ve ürün adı kullanmayın."],
        anahtar=None,
    ),
    dict(
        no="ÖZET", ad="ÖZET ve ANAHTAR KELİMELER", hedef=0, tur="ozet",
        tubitak=["\"Araştırma önerisi özetinin (1) bilimsel nitelik; (2) yöntem; (3) proje yönetimi ve "
                 "(4) yaygın etki hakkında bilgileri kapsaması beklenir. Bu bölümün en son yazılması "
                 "önerilir.\"",
                 "En az 25, en fazla 450 kelime. Anahtar kelime: en az 3, en fazla 5, aralarında virgül."],
        panelist=["Panelist 40 öneri okuyor. Özet, ilk izlenimi kuran tek paragraftır.",
                  "Özetin dört rubrik bloğunu da AYNI SIRAYLA içermesi, panelistin puanlama formunu "
                  "doldururken aradığı şeyi hemen bulmasını sağlar.",
                  "Kelime sayısını Word'de SAYIN. 450'yi aşarsanız ön incelemede eleme riski doğar; "
                  "biz 380–445 aralığını hedefliyoruz (güvenlik payı).",
                  "Anahtar kelimeleri arama yapılabilir terimler seçin; proje adınızı anahtar kelime yapmayın."],
        anahtar="ozet",
    ),
    dict(
        no="1.1", ad="Konunun Önemi ve Araştırma Önerisinin Bilimsel Niteliği",
        hedef=2200, tur="metin",
        tubitak=["\"Araştırma önerisinde ele alınan konunun kapsamı, sınırları ve önemi ortaya konulur. "
                 "Araştırma önerisi kapsamında yapılacak çalışmalarla literatürdeki hangi eksikliğin nasıl "
                 "giderileceği veya hangi soruna nasıl bir çözüm getirileceği ilgili literatüre atıfla "
                 "açıklanarak araştırma önerisinin bilimsel niteliği ortaya konulur. Araştırma sorusu ve "
                 "varsa hipotez(ler)i tanımlanır.\"",
                 "\"Projenin konusu, 12. Kalkınma Planı ve 2030 Sanayi ve Teknoloji Stratejisi'nde yer alan "
                 "kritik teknoloji alanları ile öncelikli Ar-Ge ve yenilik konuları ile ilişkili ise, "
                 "ilişkilendirilme sebebi ve ilgili alana sağlayacağı yararlar açıklanmalıdır.\"",
                 "En az 350 karakter."],
        panelist=["AĞIRLIK: %35 bloğunun kalbi. Buradan kaybettiğiniz puanı başka yerden telafi edemezsiniz.",
                  "Panelistin aradığı zincir: PROBLEM → LİTERATÜR → BOŞLUK → ARAŞTIRMA SORUSU → HİPOTEZ. "
                  "Zincirin bir halkası kopuksa puan düşer.",
                  "En sık ret gerekçesi: araştırma sorusu yerine yapılacaklar listesi. \"Şunu geliştireceğiz\" "
                  "değil, \"şu değişken şu çıktıyı ne yönde etkiler?\" yazın.",
                  "\"Literatürde çalışma azdır\" cümlesi atıfsız yazılırsa iddia sayılmaz. Boşluğu şu kalıpla "
                  "kanıtlayın: \"[3] X'i ele almış ancak Y koşulunu incelememiştir; [7] Y'yi incelemiş ancak "
                  "Z dilinde sınamamıştır.\"",
                  "Madde 5.2.4 avantajı: 12. Kalkınma Planı / 2030 Stratejisi ilişkisi kurulan başvurular için "
                  "destek taban puanı AYRI belirlenebilir. Bu paragrafı atlamayın — bedava avantaj."],
        anahtar="bolum_1_1",
    ),
    dict(
        no="1.2", ad="Amaç ve Hedefler", hedef=1400, tur="metin",
        tubitak=["\"Araştırma önerisinin amacı ve hedefleri açık, ölçülebilir, gerçekçi ve ulaşılabilir "
                 "nitelikte olacak şekilde yazılır.\"",
                 "En az 350 karakter."],
        panelist=["AĞIRLIK: %35 bloğunun ikinci yarısı.",
                  "\"Ölçülebilir\" kelimesi formun kendi sözü. Her özgül hedefin yanında SAYISAL bir eşik "
                  "olmalı: \"makro-F1 ≥ 0,85\", \"uzman uyumu κ ≥ 0,70\", \"hata oranında ≥ %20 azalma\".",
                  "Tek genel amaç + 5–7 numaralı özgül hedef en okunaklı yapıdır.",
                  "Her hedefi bir iş paketiyle eşleyin (Ö1 → İP2 gibi). Panelist bu eşlemeyi arar; yoksa "
                  "proje yönetimi puanı da düşer.",
                  "\"Ulaşılabilir\" kelimesi bir uyarıdır: 12 ayda ve 12.000 TL ile yapılamayacak bir hedef "
                  "listesi güvenilirliği düşürür."],
        anahtar="bolum_1_2",
    ),
    dict(
        no="2", ad="YÖNTEM", hedef=3000, tur="metin",
        tubitak=["\"Araştırmada uygulanacak yöntem ve araştırma tekniklerinin, amaç ve hedeflere ulaşmaya "
                 "ne düzeyde elverişli olduğu ilgili literatüre atıf yapılarak ortaya konulur.\"",
                 "\"Yöntem bölümünün; araştırma tasarımı, bağımlı ve bağımsız değişkenler, istatistiksel "
                 "yöntemler vb. unsurları içermesi gerekir. Araştırma önerisinde herhangi bir ön çalışma "
                 "veya fizibilite yapıldıysa bunların sunulması beklenir. Araştırma önerisinde sunulan "
                 "yöntemlerin çalışma takvimi ile ilişkilendirilmesi gerekir.\"",
                 "En az 350 karakter."],
        panelist=["AĞIRLIK: %25.",
                  "Form üç şeyi ADIYLA istiyor: araştırma tasarımı, bağımlı/bağımsız değişkenler, "
                  "istatistiksel yöntemler. Üçünü de ayrı alt başlık yapın — panelist tik atarak okur.",
                  "Bağımsız değişkeni gerçekten MANİPÜLE ettiğinizi gösterin. \"Sistemi yaptık ve ölçtük\" "
                  "kontrollü deney değildir; en az bir karşılaştırma temeli (baseline) ve bir ablasyon koşulu olmalı.",
                  "Her hipotezin karşısına uygulanacak testi yazın. Güç analizi (etki büyüklüğü, α, güç, "
                  "gereken örneklem) yazılmamışsa örneklem sayınız keyfî görünür.",
                  "\"Ön çalışma yapıldıysa sunulması BEKLENİR\" — mevcut prototipiniz burada AVANTAJDIR. "
                  "Gizlemeyin; ölçümleriyle sunun. Eski/güvenilmez ölçümleri ise 'proje kapsamında yeniden "
                  "ölçülecektir' diye işaretleyin.",
                  "Kişisel veri kullanılıp kullanılmadığını ve etik kurul durumunu AÇIKÇA yazın. "
                  "Belirsiz bırakılırsa proje yürürlüğe alınırken tıkanır (m.5.3.1)."],
        anahtar="bolum_2",
    ),
    dict(
        no="3.1", ad="Çalışma Takvimi", hedef=0, tur="takvim",
        tubitak=["\"Araştırmada yer alacak başlıca faaliyetler, her bir faaliyetin hangi sürede "
                 "gerçekleştirileceği, başarı ölçütü ve araştırmanın başarısına katkısı 'Çalışma Takvimi' "
                 "doldurularak sunulur.\"",
                 "Dipnot (**): \"Literatür taraması, sonuç raporu hazırlama aşamaları, araştırma "
                 "sonuçlarının paylaşımı ve malzeme alımı ayrı birer iş adımı olarak gösterilmemelidir.\"",
                 "Dipnot (***): \"Başarı ölçütü, ölçülebilir ve izlenebilir nitelikte olacak şekilde nicel "
                 "veya nitel ölçütlerle (ifade, sayı, yüzde, vb.) belirtilir. Bu sütundaki değerlerin "
                 "TOPLAMI 100 OLMALIDIR.\""],
        panelist=["AĞIRLIK: %20 bloğunun en görünür parçası.",
                  "Katkı sütununun toplamı 99 veya 101 ise panelist bunu dikkatsizlik olarak okur. TOPLAYIN.",
                  "Dipnot (**) yasağı en sık düşülen tuzaktır: literatür taramasını ayrı satır yapmayın, "
                  "ilgili iş paketinin faaliyetleri içine gömün.",
                  "Başarı ölçütü \"tamamlanmış olması\" değildir. \"Ölçülebilir ve izlenebilir\" kelimesi "
                  "formun kendi şartıdır: sayı, yüzde veya eşik verin.",
                  "12 ayı aşmayın (m.2.2/m.6.1). İlk ay veri/altyapı, son ay değerlendirme ve raporlama "
                  "içerecek şekilde dengeleyin."],
        anahtar="calisma_takvimi",
    ),
    dict(
        no="3.2", ad="Risk Yönetimi", hedef=0, tur="risk",
        tubitak=["\"Araştırmanın başarısını olumsuz yönde etkileyebilecek riskler ve bu risklerle "
                 "karşılaşıldığında araştırmanın başarıyla yürütülmesini sağlamak için alınacak tedbirler "
                 "(B Planı) Risk Yönetimi Tablosu'nda ifade edilir. B Plan(lar)ının uygulanması araştırmanın "
                 "temel hedeflerinden sapmaya yol açmamalıdır. B Plan(lar)ına geçilmesi durumunda yöntem "
                 "değişikliğine gidiliyor ise bu durum detaylandırılmalıdır.\""],
        panelist=["\"Risk yok\" veya \"proje gecikebilir\" yazmak puan kaybettirir.",
                  "İyi risk somuttur ve projeye özgüdür: veri erişimi gecikmesi, etiketleyici uyumunun düşük "
                  "çıkması, örneklem yetersizliği, ekip üyesi ayrılması, dış hizmetin fiyat/erişim değişimi.",
                  "Formun altını çizdiği şart: B planı ana hedeften SAPTIRMAMALI. \"Yapamazsak kapsamı "
                  "daraltırız\" cevabı bu şartı ihlal eder.",
                  "4–6 risk yeterlidir. Her biri için kim, ne zaman, neye göre devreye alacak — yazın."],
        anahtar="riskler",
    ),
    dict(
        no="3.3", ad="Araştırma Olanakları", hedef=0, tur="olanak",
        tubitak=["\"Bu bölümde projenin yürütüleceği kurum ve kuruluşlarda var olan ve projede kullanılacak "
                 "olan altyapı/ekipman (laboratuvar, araç, makine-teçhizat, vb.) olanakları belirtilir.\""],
        panelist=["Danışmanınız başvuruyu onaylarken \"bölüm olanaklarının yeterli olduğunu\" taahhüt eder "
                  "(m.4.3.4). Bu tabloyla o taahhüt tutarlı olmalı.",
                  "Bütçede istediğiniz bir ekipman burada \"zaten var\" diye yazılıysa çelişki doğar — "
                  "bütçeyle bu tabloyu karşılaştırın.",
                  "Yazılım/hesaplama kaynakları da altyapıdır; marka adı vermeden türüyle yazın."],
        anahtar="olanaklar",
    ),
    dict(
        no="4", ad="ARAŞTIRMA ÖNERİSİNİN YAYGIN ETKİSİ", hedef=1800, tur="etki",
        tubitak=["\"Araştırma önerisi kapsamındaki çalışmadan elde edilmesi öngörülen çıktılar amaçlarına "
                 "göre belirlenen kategorilere ayrılarak belirtilir; ölçülebilir ve gerçekçi hedeflere "
                 "dayandırılır.\"",
                 "Üç kategori: (1) Bilimsel/Akademik Çıktılar, (2) Ekonomik/Ticari/Sosyal Çıktılar, "
                 "(3) Yeni Proje(ler) Oluşturmasına Yönelik Çıktılar."],
        panelist=["AĞIRLIK: %20.",
                  "Anahtar kelime \"gerçekçi\". Lisans projesinde 5 SCI makalesi vaadi güvenilirliği düşürür; "
                  "1–2 ulusal bildiri + 1 makale hedefi inandırıcıdır.",
                  "Her çıktıya SAYI ve ZAMAN verin: \"proje bitiminden itibaren 6 ay içinde 1 ulusal "
                  "hakemli dergi makalesi\".",
                  "Açık veri kümesi / açık kaynak yayımlama taahhüdü somut ve ölçülebilir bir çıktıdır — "
                  "ekleyin.",
                  "Madde 2.1.2 hatırlatması: konferans katılım, yayın ve patent MASRAFLARI bütçeye "
                  "konulamaz. Ama bunları ÇIKTI olarak hedeflemek serbesttir.",
                  "Madde 7'ye köprü kurun: proje tamamlanınca 2224-B gibi programlara başvuru imkânı doğar."],
        anahtar="yaygin_etki",
    ),
    dict(
        no="5", ad="BELİRTMEK İSTEDİĞİNİZ DİĞER KONULAR", hedef=0, tur="metin",
        tubitak=["\"Sadece araştırma önerisinin değerlendirilmesine katkı sağlayabilecek bilgi/veri "
                 "(grafik, tablo, vb.) eklenebilir.\""],
        panelist=["Boş bırakmak kayıp değildir; ama iyi kullanılırsa kazançtır.",
                  "Buraya konacak en güçlü şeyler: ön çalışma sonuç grafiği, sistem/akış şeması, "
                  "deney tasarımı tablosu, ekip yetkinlik özeti.",
                  "Sayfa sınırını (20) unutmayın — buradaki her görsel sayfa harcar."],
        anahtar="bolum_5_diger",
    ),
    dict(
        no="6", ad="EKLER — EK-1: KAYNAKLAR", hedef=0, tur="kaynak",
        tubitak=["Form EK-1 olarak KAYNAKLAR bölümünü ister. Ekler 20 sayfalık sınıra DÂHİL DEĞİLDİR."],
        panelist=["Metindeki her [n] atıfının kaynakçada karşılığı olmalı; kaynakçadaki her kayıt metinde "
                  "en az bir kez kullanılmalı. Panelist bunu rastgele örnekleyerek kontrol eder.",
                  "Güncellik önemli: kaynakların önemli bir bölümü son 5 yıldan olsun.",
                  "Var olmayan kaynak yazmak etik ihlalidir ve AYEK Yönetmeliği kapsamında işlem yapılır "
                  "(m.5.2.3, m.8.2). ÜYZ araçlarının ürettiği her atfı tek tek doğrulayın.",
                  "Tek bir atıf biçimi kullanın (APA 7 önerilir) ve sonuna kadar tutarlı kalın."],
        anahtar="kaynaklar",
    ),
]


def uret(icerik, cikti):
    doc = Document()
    sayfa_kur(doc)

    p = doc.add_paragraph()
    _font(p.add_run("TÜBİTAK 2209-A"), size=20, bold=True, renk=RGBColor(0x0D, 0x2B, 0x45))
    p = doc.add_paragraph()
    _font(p.add_run("Araştırma Önerisi Formu — REHBERLİ ŞABLON"), size=13, bold=True,
          renk=RGBColor(0x8B, 0x25, 0x00))
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    _font(p.add_run("Her bölümde dört katman: TÜBİTAK'ın resmî açıklaması · panelistin ne aradığı · "
                    "doldurulmuş örnek · sizin yazacağınız alan"), size=10,
          renk=RGBColor(0x44, 0x44, 0x44))

    kutu(doc, "BU BELGE NASIL KULLANILIR", [
        ("1. Bu dosya BAŞVURUYA YÜKLENMEZ.", {"bold": True}),
        "Rehber kutuları çalışma amaçlıdır ve 20 sayfa sınırını aşar. Başvuruya yükleyeceğiniz dosya, "
        "aynı klasördeki TESLİM dosyasıdır (resmî TÜBİTAK formunun doldurulmuş hâli).",
        ("2. Renk kodları:", {"bold": True}),
        ("MAVİ kutu = TÜBİTAK'ın formdaki/çağrı metnindeki kendi ifadesi (birebir alıntı).", {"bullet": True}),
        ("KİREMİT kutu = panelistin ne aradığı, puan ağırlığı ve puan kesme riski.", {"bullet": True}),
        ("YEŞİL kutu = bu projeye göre doldurulmuş örnek metin — kopyalayıp uyarlayın.", {"bullet": True}),
        ("GRİ kutu = sizin yazacağınız alan; karakter hedefi üstünde yazar.", {"bullet": True}),
        ("3. Yazdıktan sonra aynı klasördeki KONTROL LİSTESİ belgesinden geçirin.", {"bold": True}),
    ], "panelist")

    doc.add_page_break()

    for b in BOLUM_REHBERI:
        # "ÖZET. ÖZET ve ANAHTAR KELİMELER" gibi tekrarı önle: bölüm kodu
        # numarasal değilse başlığın kendisi zaten kodu içeriyor demektir.
        no = b["no"]
        baslik(doc, f"{no}. {b['ad']}" if re.match(r"^[\dA-Z]$|^\d+\.\d+$", no) else b["ad"])
        kutu(doc, "TÜBİTAK'IN RESMÎ AÇIKLAMASI", b["tubitak"], "tubitak")
        kutu(doc, "PANELİST NE ARAR", [(s, {"bullet": True}) for s in b["panelist"]], "panelist")

        anahtar = b["anahtar"]
        if anahtar:
            deger = icerik.get(anahtar)
            if b["tur"] == "ozet":
                satirlar = alinti(icerik.get("ozet", ""), 1600)
                ak = icerik.get("anahtar_kelimeler", [])
                if ak:
                    satirlar.append(("Anahtar Kelimeler: " + ", ".join(ak), {"bold": True}))
                kutu(doc, "ÖRNEK — DOLDURULMUŞ", satirlar, "ornek")
            elif b["tur"] == "takvim" and deger:
                kutu(doc, "ÖRNEK — DOLDURULMUŞ İŞ PAKETLERİ", [
                    (f"{ip.get('tarih_araligi','')} · {ip.get('ad','')} — "
                     f"Başarı ölçütü: {ip.get('basari_olcutu','')} (katkı %{ip.get('katki','')})",
                     {"bullet": True}) for ip in deger
                ] + [(f"KATKI TOPLAMI: %{sum(ip.get('katki', 0) for ip in deger)} "
                      f"— tam 100 olmalı", {"bold": True})], "ornek")
            elif b["tur"] == "risk" and deger:
                kutu(doc, "ÖRNEK — DOLDURULMUŞ RİSKLER", [
                    (f"RİSK: {r.get('risk','')}  →  B PLANI: {r.get('tedbir','')}", {"bullet": True})
                    for r in deger], "ornek")
            elif b["tur"] == "olanak" and deger:
                kutu(doc, "ÖRNEK — DOLDURULMUŞ OLANAKLAR", [
                    (f"{o.get('altyapi','')}  →  {o.get('kullanim_amaci','')}", {"bullet": True})
                    for o in deger], "ornek")
            elif b["tur"] == "etki" and deger:
                kutu(doc, "ÖRNEK — DOLDURULMUŞ YAYGIN ETKİ", [
                    ("Bilimsel/Akademik:", {"bold": True}), deger.get("bilimsel_akademik", ""),
                    ("Ekonomik/Ticari/Sosyal:", {"bold": True}), deger.get("ekonomik_ticari_sosyal", ""),
                    ("Yeni Projeler:", {"bold": True}), deger.get("yeni_projeler", ""),
                ], "ornek")
            elif b["tur"] == "kaynak" and deger:
                kutu(doc, f"ÖRNEK — KAYNAKÇA ({len(deger)} kayıt, ilk 8'i gösteriliyor)",
                     [(k, {"size": 8}) for k in deger[:8]] +
                     [("[...] Tam liste TESLİM dosyasının EK-1 bölümündedir.", {"italic": True})], "ornek")
            elif isinstance(deger, str) and deger.strip():
                kutu(doc, "ÖRNEK — DOLDURULMUŞ", alinti(deger), "ornek")

        if b["tur"] == "ozet":
            bos_alan(doc, 0,
                     "İpucu: en son yazın. Dört rubrik bloğunu aynı sırayla geçirin: "
                     "bilimsel nitelik → yöntem → proje yönetimi → yaygın etki.",
                     etiket="▼ BURAYA SİZ YAZIN — 380–445 kelime hedefleyin "
                            "(TÜBİTAK sınırı 25–450; Word > Sözcük Sayımı ile SAYIN)",
                     satir=8)
            bos_alan(doc, 0, "", etiket="▼ ANAHTAR KELİMELER — tam 5 adet, aralarına virgül", satir=1)
        elif b["hedef"]:
            bos_alan(doc, b["hedef"],
                     "İpucu: örnek metnin YAPISINI koruyun, içeriği kendi projenize göre değiştirin.")
        doc.add_page_break()

    # kapanis
    baslik(doc, "SON KONTROL")
    kutu(doc, "GÖNDERMEDEN ÖNCE", [
        ("Bu şablonu değil, TESLİM dosyasını yükleyin.", {"bullet": True}),
        ("KONTROL LİSTESİ belgesindeki ön inceleme kapısındaki her kutuyu işaretleyin.", {"bullet": True}),
        ("Word'de sayfa sayısını sayın: ekler hariç 20'yi geçmemeli.", {"bullet": True}),
        ("Özet kelime sayısını Word > Sözcük Sayımı ile doğrulayın (25–450).", {"bullet": True}),
        ("Çalışma takvimi katkı sütununu toplayın: tam 100.", {"bullet": True}),
        ("Metinde marka/kurum/ürün adı arayın (Ctrl+F): sıfır olmalı.", {"bullet": True}),
        ("Her kaynağın gerçekten var olduğunu tek tek doğrulayın.", {"bullet": True}),
        ("Danışman onayı + kuruluş yetkilisi e-imzası için en az 1 hafta pay bırakın.", {"bullet": True}),
    ], "uyari")

    doc.save(cikti)
    return cikti


if __name__ == "__main__":
    with open(sys.argv[1], encoding="utf-8") as f:
        icerik = json.load(f)
    print("YAZILDI:", uret(icerik, sys.argv[2]))
