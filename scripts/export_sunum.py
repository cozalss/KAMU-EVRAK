#!/usr/bin/env python
"""Sunum Markdown → PDF + PPTX export (TEKNOFEST teslim formatı, şartname §7).

Anlatım Markdown'ları (`docs/sunum-on-degerlendirme.md`, `docs/sunum-final.md`)
gerçek **PDF** ve **PPTX** dosyalarına dönüştürür. Saf-python: PDF için çekirdek
`reportlab` (TTF font gömme → Türkçe ç/ş/ğ/İ DOĞRU), PPTX için `python-pptx`.
Sistem bağımlılığı (pandoc/LibreOffice) YOK.

    # Tek sunum → docs/sunum/on-degerlendirme.pdf + .pptx
    python scripts/export_sunum.py docs/sunum-on-degerlendirme.md \
        --out docs/sunum/on-degerlendirme
    # Her iki sunumu da üret (varsayılan çıktı docs/sunum/):
    python scripts/export_sunum.py --all

Parse kuralı (bkz. docs/sunum/README.md):
  * `## Slayt N — Başlık`  → yeni slayt başlığı (`## ` herhangi bir başlık de olur)
  * `- madde` / `* madde`  → bullet (girinti = alt-bullet seviyesi)
  * `**Etiket:** değer`    → bullet (kalın etiket düz metne indirgenir)
  * `> Konuşma notu ...`   → speaker notes (PPTX notes; PDF'te küçük italik blok)
  * `![alt](docs/img/x.png)` veya satır içi `docs/img/*.png` → ilgili görsel gömülür
  * Markdown tablosu (`| a | b |`) → satırlar bullet'a indirgenir (basit, okunur)

Bu script SADECE okur + üretir; backend graph/node/scorecard'a dokunmaz.
"""

from __future__ import annotations

import argparse
import contextlib
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

with contextlib.suppress(Exception):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")  # type: ignore[union-attr]

_REPO = Path(__file__).resolve().parent.parent
_DEFAULT_OUT_DIR = _REPO / "docs" / "sunum"
_ALL_SUNUMLAR: tuple[tuple[str, str], ...] = (
    ("docs/sunum-on-degerlendirme.md", "on-degerlendirme"),
    ("docs/sunum-final.md", "final"),
)

# Türkçe-yetenekli TTF adayları (Windows Arial/Tahoma/Verdana hepsi ç/ş/ğ/İ kapsar).
# İlk bulunan (regular, bold) çifti gömülür; hiçbiri yoksa NET hata verilir.
_FONT_ADAYLARI: tuple[tuple[str, str, str], ...] = (
    ("Arial", "C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
    ("DejaVuSans", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
     "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ("Verdana", "C:/Windows/Fonts/verdana.ttf", "C:/Windows/Fonts/verdanab.ttf"),
    ("Tahoma", "C:/Windows/Fonts/tahoma.ttf", "C:/Windows/Fonts/tahomabd.ttf"),
)

# --- Parse regex'leri (modül-seviyesi: tek-derleme) ---
_RE_SLIDE = re.compile(r"^##\s+(.*)$")
_RE_H1 = re.compile(r"^#\s+(.*)$")
_RE_BULLET = re.compile(r"^(\s*)[-*]\s+(.*)$")
_RE_NOTE = re.compile(r"^>\s?(.*)$")
_RE_IMG = re.compile(r"!\[[^\]]*\]\(([^)]+)\)")
_RE_IMG_BARE = re.compile(r"`?(docs/img/[\w./-]+\.(?:png|jpg|jpeg|gif))`?")
_RE_TABLE_ROW = re.compile(r"^\s*\|(.+)\|\s*$")
_RE_TABLE_SEP = re.compile(r"^\s*\|[\s:|-]+\|\s*$")
_RE_BOLD = re.compile(r"\*\*([^*]+)\*\*")
_RE_INLINE_CODE = re.compile(r"`([^`]+)`")
_RE_LINK = re.compile(r"\[([^\]]+)\]\([^)]+\)")


@dataclass
class Bullet:
    """Tek bir madde; level=0 ana, level>=1 alt-bullet."""

    text: str
    level: int = 0


@dataclass
class Slide:
    """Parse edilmiş tek slayt: başlık + bulletlar + notlar + görsel(ler)."""

    title: str
    bullets: list[Bullet] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)
    images: list[str] = field(default_factory=list)


def _strip_inline(text: str) -> str:
    """Markdown satır-içi işaretlerini düz metne indir (kalın/kod/link)."""
    text = _RE_BOLD.sub(r"\1", text)
    text = _RE_INLINE_CODE.sub(r"\1", text)
    text = _RE_LINK.sub(r"\1", text)
    return text.replace(" ", " ").strip()


def _find_images(text: str) -> list[str]:
    """Bir satırdaki tüm görsel yollarını (md veya çıplak) repo-göreli döndür."""
    found = list(_RE_IMG.findall(text))
    for bare in _RE_IMG_BARE.findall(text):
        if bare not in found:
            found.append(bare)
    return found


def _table_to_bullets(rows: list[str]) -> list[Bullet]:
    """Markdown tablosunu okunur bulletlara indir: başlık + 'sütun: değer' satırları."""
    parsed = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    if not parsed:
        return []
    header, *body = parsed
    out: list[Bullet] = []
    for cells in body:
        primary = _strip_inline(cells[0]) if cells else ""
        rest = [
            f"{_strip_inline(header[i])}: {_strip_inline(c)}"
            for i, c in enumerate(cells[1:], start=1)
            if i < len(header) and _strip_inline(c)
        ]
        out.append(Bullet(primary, 0))
        out.extend(Bullet(r, 1) for r in rest)
    return out


def _flush_table(buf: list[str], slide: Slide | None) -> None:
    """Biriken tablo satırlarını slayt bulletlarına ekle ve tamponu boşalt."""
    if buf and slide is not None:
        slide.bullets.extend(_table_to_bullets(buf))
    buf.clear()


def parse_markdown(md_text: str) -> tuple[str, list[Slide]]:
    """Sunum Markdown'ını (deck-başlığı, slaytlar) olarak ayrıştır."""
    deck_title = "Sunum"
    slides: list[Slide] = []
    current: Slide | None = None
    table_buf: list[str] = []
    in_code = False

    for raw in md_text.splitlines():
        line = raw.rstrip("\n")
        if line.strip().startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            continue

        if (m := _RE_H1.match(line)) and not slides and current is None:
            deck_title = _strip_inline(m.group(1))
            continue
        if m := _RE_SLIDE.match(line):
            _flush_table(table_buf, current)
            current = Slide(title=_strip_inline(m.group(1)))
            slides.append(current)
            continue
        if current is None:
            continue
        _consume_body_line(line, current, table_buf)

    _flush_table(table_buf, current)
    return deck_title, slides


def _consume_body_line(line: str, slide: Slide, table_buf: list[str]) -> None:
    """Slayt gövdesindeki tek satırı uygun alana (bullet/not/görsel/tablo) yerleştir."""
    if _RE_TABLE_SEP.match(line):
        return
    if _RE_TABLE_ROW.match(line):
        table_buf.append(line)
        return
    _flush_table(table_buf, slide)

    for img in _find_images(line):
        if img not in slide.images:
            slide.images.append(img)

    if m := _RE_NOTE.match(line):
        note = _strip_inline(m.group(1))
        if note:
            slide.notes.append(note)
        return
    if m := _RE_BULLET.match(line):
        indent = len(m.group(1).replace("\t", "  "))
        text = _strip_inline(m.group(2))
        if text:
            slide.bullets.append(Bullet(text, level=1 if indent >= 2 else 0))
        return
    stripped = _strip_inline(line)
    if stripped and not stripped.startswith("---") and not _RE_IMG.match(line):
        slide.bullets.append(Bullet(stripped, level=0))


def _resolve_image(path_str: str) -> Path | None:
    """Repo-göreli görsel yolunu mutlak yola çevir; yoksa None."""
    candidate = (_REPO / path_str).resolve()
    return candidate if candidate.is_file() else None


def _register_font() -> tuple[str, str]:
    """İlk bulunan Türkçe-yetenekli TTF (regular, bold) çiftini reportlab'a gömer.

    Döner: (regular_font_adı, bold_font_adı). Hiçbiri yoksa RuntimeError.
    """
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    for name, reg_path, bold_path in _FONT_ADAYLARI:
        if not Path(reg_path).is_file():
            continue
        bold_name = f"{name}-Bold"
        pdfmetrics.registerFont(TTFont(name, reg_path))
        if Path(bold_path).is_file():
            pdfmetrics.registerFont(TTFont(bold_name, bold_path))
        else:
            pdfmetrics.registerFont(TTFont(bold_name, reg_path))
        with contextlib.suppress(Exception):
            pdfmetrics.registerFontFamily(name, normal=name, bold=bold_name)
        return name, bold_name
    raise RuntimeError(
        "Türkçe-yetenekli TTF bulunamadı (Arial/DejaVu/Verdana/Tahoma). "
        "Bir TTF kurun veya _FONT_ADAYLARI'na yol ekleyin."
    )


def _pdf_styles(font: str, bold: str):
    """PDF paragraf stilleri (başlık/bullet/alt-bullet/not) — Türkçe font gömülü."""
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.styles import ParagraphStyle

    return {
        "title": ParagraphStyle(
            "Title", fontName=bold, fontSize=20, leading=24,
            spaceAfter=14, textColor="#16324f",
        ),
        "bullet": ParagraphStyle(
            "Bullet", fontName=font, fontSize=13, leading=18,
            leftIndent=16, bulletIndent=4, spaceAfter=6, alignment=TA_LEFT,
        ),
        "subbullet": ParagraphStyle(
            "SubBullet", fontName=font, fontSize=11.5, leading=16,
            leftIndent=34, bulletIndent=22, spaceAfter=4, textColor="#333333",
        ),
        "note": ParagraphStyle(
            "Note", fontName=font, fontSize=9.5, leading=13,
            leftIndent=12, spaceBefore=8, textColor="#5a6a7a",
        ),
    }


def _pdf_slide_flowables(slide: Slide, idx: int, styles: dict) -> list:
    """Tek slaytın PDF flowable listesi (başlık + görsel + bullet + not)."""
    from reportlab.lib.units import cm
    from reportlab.platypus import Image, PageBreak, Paragraph, Spacer
    from xml.sax.saxutils import escape

    flow: list = []
    if idx > 0:
        flow.append(PageBreak())
    flow.append(Paragraph(escape(slide.title), styles["title"]))
    for img in slide.images:
        if resolved := _resolve_image(img):
            with contextlib.suppress(Exception):
                pic = Image(str(resolved))
                pic._restrictSize(16 * cm, 7.5 * cm)
                flow.extend([pic, Spacer(1, 0.3 * cm)])
    for b in slide.bullets:
        style = styles["subbullet"] if b.level else styles["bullet"]
        marker = "◦" if b.level else "•"
        flow.append(Paragraph(f"{marker}&nbsp;{escape(b.text)}", style))
    for note in slide.notes:
        flow.append(Paragraph(f"<i>Not: {escape(note)}</i>", styles["note"]))
    return flow


def build_pdf(deck_title: str, slides: list[Slide], out_path: Path) -> Path:
    """Slaytları tek PDF'e (slayt-başına-sayfa) render et; Türkçe TTF gömülü."""
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import SimpleDocTemplate

    font, bold = _register_font()
    styles = _pdf_styles(font, bold)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(
        str(out_path), pagesize=landscape(A4),
        title=deck_title, leftMargin=42, rightMargin=42, topMargin=36, bottomMargin=36,
    )
    story: list = []
    for idx, slide in enumerate(slides):
        story.extend(_pdf_slide_flowables(slide, idx, styles))
    doc.build(story)
    return out_path


def _pptx_fill_bullets(body, bullets: list[Bullet]) -> None:
    """Bir içerik placeholder'ına bullet paragraflarını (seviyeli) yaz."""
    tf = body.text_frame
    tf.word_wrap = True
    first = True
    for b in bullets:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        para.text = b.text
        para.level = min(b.level, 4)
        first = False
    if first:  # hiç bullet yoksa boş placeholder bırakma
        tf.paragraphs[0].text = ""


def _pptx_add_image(slide, images: list[str], prs) -> None:
    """İlk çözülebilen görseli slaytın sağ-alt köşesine ölçekli ekle."""
    from pptx.util import Emu

    for img in images:
        resolved = _resolve_image(img)
        if not resolved:
            continue
        with contextlib.suppress(Exception):
            width = Emu(int(prs.slide_width * 0.42))
            left = prs.slide_width - width - Emu(int(prs.slide_width * 0.03))
            top = Emu(int(prs.slide_height * 0.30))
            slide.shapes.add_picture(str(resolved), left, top, width=width)
            return


def build_pptx(deck_title: str, slides: list[Slide], out_path: Path) -> Path:
    """Slaytları PPTX'e render et: başlık + bullet + görsel + speaker notes."""
    from pptx import Presentation

    prs = Presentation()
    bullet_layout = prs.slide_layouts[1]  # "Title and Content"
    for slide_data in slides:
        slide = prs.slides.add_slide(bullet_layout)
        slide.shapes.title.text = slide_data.title or deck_title
        body = slide.placeholders[1]
        _pptx_fill_bullets(body, slide_data.bullets)
        _pptx_add_image(slide, slide_data.images, prs)
        if slide_data.notes:
            slide.notes_slide.notes_text_frame.text = "\n".join(slide_data.notes)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return out_path


def export_one(md_path: Path, out_base: Path) -> tuple[Path, Path]:
    """Tek Markdown'ı PDF + PPTX'e dönüştür; (pdf_yolu, pptx_yolu) döner."""
    md_text = md_path.read_text(encoding="utf-8")
    deck_title, slides = parse_markdown(md_text)
    if not slides:
        raise ValueError(f"Slayt bulunamadı (## başlık yok): {md_path}")
    pdf = build_pdf(deck_title, slides, out_base.with_suffix(".pdf"))
    pptx = build_pptx(deck_title, slides, out_base.with_suffix(".pptx"))
    print(f"  ✓ {len(slides)} slayt → {pdf.name} + {pptx.name}")
    return pdf, pptx


def _parse_args(argv: list[str] | None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Sunum Markdown → PDF + PPTX (TEKNOFEST teslim formatı).",
    )
    parser.add_argument("markdown", nargs="?", help="Kaynak .md (örn. docs/sunum-final.md)")
    parser.add_argument(
        "--out", help="Çıktı taban yolu (uzantısız; .pdf/.pptx eklenir).",
    )
    parser.add_argument(
        "--all", action="store_true",
        help="Her iki sunumu da docs/sunum/ altına üret.",
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = _parse_args(argv)
    if args.all or not args.markdown:
        print("Tüm sunumlar üretiliyor → docs/sunum/")
        for rel, stem in _ALL_SUNUMLAR:
            src = _REPO / rel
            if not src.is_file():
                print(f"  ! atlandı (yok): {rel}")
                continue
            print(f"- {rel}")
            export_one(src, _DEFAULT_OUT_DIR / stem)
        return 0

    md_path = (_REPO / args.markdown).resolve() if not Path(args.markdown).is_absolute() \
        else Path(args.markdown)
    if not md_path.is_file():
        print(f"HATA: kaynak bulunamadı: {md_path}", file=sys.stderr)
        return 2
    out_base = Path(args.out).resolve() if args.out else _DEFAULT_OUT_DIR / md_path.stem
    print(f"- {md_path.name} → {out_base}")
    export_one(md_path, out_base)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
