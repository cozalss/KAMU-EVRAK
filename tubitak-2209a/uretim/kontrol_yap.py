"""C. 2209A_KONTROL_LISTESI.docx uretir.

Icerik kaynagi: TUBITAK BIDEB 2209-A 2025 Yili Cagri Duyurusu (V2.1) ve
resmi Arastirma Onerisi Formu (09.10.2025 surumu). Uydurma kural yoktur;
her madde cagri metnindeki numarali hukme atifla verilir.
"""
import sys, os
sys.stdout.reconfigure(encoding="utf-8")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from ekler import sayfa_kur, baslik, gövde, kutu, tablo_yap, _font, RENK

KUTU = "☐"


def uret(cikti):
    doc = Document()
    sayfa_kur(doc)

    # ---------------------------------------------------------------- kapak
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    _font(p.add_run("TÜBİTAK 2209-A"), size=20, bold=True,
          renk=RGBColor(0x0D, 0x2B, 0x45))
    p = doc.add_paragraph()
    _font(p.add_run("Üniversite Öğrencileri Araştırma Projeleri Destekleme Programı"),
          size=11, renk=RGBColor(0x33, 0x33, 0x33))
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    _font(p.add_run("BAŞVURU KONTROL LİSTESİ, RUBRİK ÖZ-PUANLAMA VE SÜREÇ REHBERİ"),
          size=13, bold=True, renk=RGBColor(0x8B, 0x25, 0x00))

    kutu(doc, "BU BELGE NEDİR", [
        "Bu belge başvuruya YÜKLENMEZ. Kendi başvurunuzu göndermeden önce üzerinden "
        "geçeceğiniz bir denetim listesidir.",
        "İçindeki her kural, TÜBİTAK BİDEB 2209-A Çağrı Duyurusu'nun numaralı maddesine "
        "veya resmî Araştırma Önerisi Formu'nun kendi metnine dayanır. Parantez içindeki "
        "madde numaraları çağrı duyurusuna aittir.",
        ("UYARI: Çağrı duyurusu her yıl güncellenir. Başvuru yapmadan önce "
         "tubitak.gov.tr üzerinden o dönemin güncel çağrı metnini ve başvuru tarihlerini "
         "mutlaka teyit edin.", {"bold": True}),
    ], "uyari")

    # -------------------------------------------- 1. On inceleme eleme kapisi
    baslik(doc, "1. ÖN İNCELEME ELEME KAPISI — bunlar sağlanmazsa puan bile almazsınız")
    kutu(doc, "ÇAĞRI METNİ MADDE 5.1.1 ve 5.1.2", [
        "\"Başvuru koşullarından herhangi birini sağlamayan, belgesi eksik olan veya istenen "
        "formatta olmayan başvurular... bilimsel değerlendirmeye tabi tutulmadan ön incelemede elenir.\"",
        "\"...şablonda belirtilen minimum karakter sınırlamalarına uyulmamış veya Türkçe dilinde "
        "hazırlanmamış olan araştırma önerileri yüklenerek yapılan başvurular ön incelemede elenir.\"",
    ], "tubitak")

    tablo_yap(doc,
        ["", "Kontrol", "Ölçüt", "Dayanak"],
        [
            [KUTU, "Araştırma önerisi TÜBİTAK'ın resmî formatında hazırlandı",
             "Resmî .docx formu kullanıldı, bölüm başlıkları değiştirilmedi", "m.5.1.2"],
            [KUTU, "Yazı tipi Arial 9", "Tüm metin Arial 9 punto", "Form başlığı"],
            [KUTU, "Ekler hariç toplam sayfa sayısı 20'yi geçmiyor",
             "Word > Dosya > Bilgi > sayfa sayısı ile SAYIN", "Form başlığı"],
            [KUTU, "Metnin tamamı Türkçe", "İngilizce cümle/başlık/özet kalıntısı yok", "m.4.2.1, m.5.1.2"],
            [KUTU, "Özet 25–450 kelime arasında", "Word > Sözcük Sayımı ile SAYIN", "Form, Özet kutusu"],
            [KUTU, "Anahtar kelime 3–5 adet, aralarında virgül", "Sayın", "Form, Özet kutusu"],
            [KUTU, "1.1 bölümü en az 350 karakter", "Karakter sayın (boşluk dâhil)", "Form kutusu"],
            [KUTU, "1.2 bölümü en az 350 karakter", "Karakter sayın", "Form kutusu"],
            [KUTU, "2. Yöntem bölümü en az 350 karakter", "Karakter sayın", "Form kutusu"],
            [KUTU, "Çalışma Takvimi 'katkı' değerleri toplamı TAM 100",
             "Sütunu toplayın; 99 veya 101 olmaz", "Form, dipnot (***)"],
            [KUTU, "Literatür taraması / sonuç raporu hazırlama / sonuçların paylaşımı / "
                   "malzeme alımı AYRI iş adımı olarak yazılmadı",
             "Takvim satırlarını tek tek okuyun", "Form, dipnot (**)"],
            [KUTU, "Hiçbir kurum, kuruluş, firma, ticari ürün veya marka öne çıkarılmadı",
             "Model adı, kütüphane adı, yarışma adı, şirket adı ARAYIN", "m.4.1.8"],
            [KUTU, "Ayrımcı, dışlayıcı veya rencide edici ifade yok", "Metni okuyun", "m.4.1.8"],
            [KUTU, "Proje başlığı araştırmanın amacını ve içeriğini yansıtıyor, bilimsel ve anlaşılır",
             "Başlıkta marka/kısaltma yok", "m.4.1.7"],
            [KUTU, "Proje süresi en fazla 12 ay", "Takvimin son ayı ≤ 12", "m.2.2, m.6.1"],
            [KUTU, "Bütçede konferans katılım, yayın veya patent gideri YOK",
             "Bütçe kalemlerini okuyun", "m.2.1.2"],
            [KUTU, "ARBİS kaydı yapıldı ve tüm bilgiler güncel", "arbis.tubitak.gov.tr", "m.4.2.3"],
            [KUTU, "Akademik danışmanın güncel YÖK kaydı var ve YÖKSİS/ARBİS bilgileri güncel",
             "Danışmanınıza teyit ettirin", "m.4.1.2, m.4.3.3"],
            [KUTU, "Aynı dönemde başka bir 2209-A/2209-B başvurusunda yer almıyorum",
             "Yürütücü VEYA ortak olarak", "m.4.1.3"],
            [KUTU, "Danışmanım bu dönem en fazla 5 projede görevli", "Danışmana sorun", "m.4.1.6"],
            [KUTU, "Proje ortağı sayısı en fazla 4", "Yürütücü hariç", "m.4.1.5"],
            [KUTU, "Açık Öğretim veya Hazırlık Sınıfı öğrencisi ekipte yürütücü/ortak değil",
             "Ekip listesini kontrol edin", "m.4.1.1 dipnot"],
            [KUTU, "Başvuru TYBS üzerinden çevrim içi yapıldı", "tybs.tubitak.gov.tr", "m.4.3.1"],
            [KUTU, "Yüklenen dosya açılıyor ve doğru içerikte", "İndirip AÇIN, kontrol edin", "m.4.3.1"],
            [KUTU, "Danışman onayı alındı", "Sistem üzerinden", "m.4.3.4"],
            [KUTU, "Kuruluş yetkilisi (Rektör/Rektör Yrd.) e-imzası tamamlandı",
             "eimza.tubitak.gov.tr — BU OLMADAN BAŞVURU GEÇERSİZ", "m.4.3.4.2"],
        ],
        genislikler=[Cm(0.8), Cm(7.4), Cm(5.6), Cm(2.2)])

    doc.add_page_break()

    # ----------------------------------------------- 2. Rubrik oz-puanlama
    baslik(doc, "2. RUBRİK ÖZ-PUANLAMA — panelistin gözüyle kendinize puan verin")
    kutu(doc, "ÇAĞRI METNİ MADDE 5.2.1 — RESMÎ AĞIRLIKLAR", [
        "Araştırma önerisinin bilimsel niteliği (%35) · Yöntem (%25) · "
        "Proje yönetimi (%20) · Araştırma önerisinin yaygın etkisi (%20)",
        "Değerlendirme YALNIZCA sisteme yüklenen belgeler üzerinden yapılır (m.5.2.2). "
        "Kafanızdaki ama kâğıda dökülmemiş hiçbir şey puan getirmez.",
    ], "tubitak")

    gövde(doc, "Her satıra 0–4 arası puan verin (0 = hiç yok, 4 = mükemmel). "
               "Ağırlıklı toplam 100 üzerindendir. 2 ve altı alan her satır, panelde "
               "puan kaybettiğiniz yerdir — önce oraları düzeltin.", size=9)

    def blok(ad, agirlik, maddeler):
        gövde(doc, f"{ad}  —  ağırlık %{agirlik}", size=10, bold=True,
              renk=RGBColor(0x8B, 0x25, 0x00))
        tablo_yap(doc, ["Değerlendirme sorusu", "Nerede kanıtlanır", "Puan (0-4)"],
                  [[m[0], m[1], ""] for m in maddeler],
                  genislikler=[Cm(9.4), Cm(5.2), Cm(1.8)])

    blok("A. ARAŞTIRMA ÖNERİSİNİN BİLİMSEL NİTELİĞİ", 35, [
        ("Ele alınan konunun kapsamı, sınırları ve önemi net biçimde ortaya konmuş mu?", "Bölüm 1.1"),
        ("Literatürdeki BOŞLUK somut olarak tanımlanmış ve atıfla gösterilmiş mi?", "Bölüm 1.1"),
        ("Araştırma sorusu açıkça yazılmış mı (soru cümlesi olarak)?", "Bölüm 1.1"),
        ("Hipotez(ler) yanlışlanabilir biçimde kurulmuş mu (H0/H1)?", "Bölüm 1.1"),
        ("Öneri, mevcut çözümlerden NEYİ farklı yapıyor — özgünlük açık mı?", "Bölüm 1.1"),
        ("12. Kalkınma Planı / 2030 Sanayi ve Teknoloji Stratejisi ilişkisi kurulmuş mu? (m.5.2.4 — "
         "destek taban puanı avantajı)", "Bölüm 1.1"),
        ("Amaç tek ve net mi; hedefler numaralı, ölçülebilir ve ulaşılabilir mi?", "Bölüm 1.2"),
        ("Her hedefin sayısal başarı eşiği var mı (ör. F1 ≥ 0,85)?", "Bölüm 1.2"),
    ])

    blok("B. YÖNTEM", 25, [
        ("Araştırma tasarımı açıkça tanımlı mı (deneysel/betimsel/karma)?", "Bölüm 2"),
        ("Bağımlı ve bağımsız değişkenler AÇIKÇA listelenmiş mi?", "Bölüm 2"),
        ("Kontrol değişkenleri ve karşılaştırma temeli (baseline) var mı?", "Bölüm 2"),
        ("İstatistiksel yöntemler adıyla belirtilmiş ve hipotezlerle eşleşmiş mi?", "Bölüm 2"),
        ("Örneklem büyüklüğü gerekçeli mi (güç analizi)?", "Bölüm 2"),
        ("Veri toplama araçları ve analiz yöntemi literatüre atıfla gerekçelendirilmiş mi?", "Bölüm 2"),
        ("Ön çalışma / fizibilite bulguları sunulmuş mu? (Form bunu açıkça istiyor)", "Bölüm 2"),
        ("Yöntem, çalışma takvimiyle ilişkilendirilmiş mi (hangi yöntem hangi İP'de)?", "Bölüm 2 + 3.1"),
        ("Veri gizliliği / kişisel veri durumu açıkça ele alınmış mı?", "Bölüm 2"),
    ])

    blok("C. PROJE YÖNETİMİ", 20, [
        ("İş paketleri mantıklı sırada ve 12 aya gerçekçi biçimde yayılmış mı?", "Bölüm 3.1"),
        ("Her iş paketinin başarı ölçütü NİCEL ve izlenebilir mi?", "Bölüm 3.1"),
        ("Katkı yüzdeleri toplamı tam 100 mü?", "Bölüm 3.1"),
        ("Kimin ne yapacağı belli mi (yürütücü / ortak / danışman rolleri)?", "Bölüm 3.1"),
        ("Riskler gerçekçi mi ve her riskin uygulanabilir bir B planı var mı?", "Bölüm 3.2"),
        ("B planları ana hedeflerden sapmaya yol açmıyor mu?", "Bölüm 3.2"),
        ("Araştırma olanakları (laboratuvar/ekipman) projeyi yürütmeye yeterli mi?", "Bölüm 3.3"),
        ("Bütçe kalemleri iş paketleriyle tutarlı mı?", "Bütçe"),
    ])

    blok("D. YAYGIN ETKİ", 20, [
        ("Bilimsel/akademik çıktılar SAYILI ve gerçekçi mi (lisans düzeyi için abartısız)?", "Bölüm 4"),
        ("Ekonomik/ticari/sosyal çıktılar somut mu (prototip, veri kümesi, çalıştay)?", "Bölüm 4"),
        ("Yeni proje(ler) oluşturma potansiyeli gösterilmiş mi?", "Bölüm 4"),
        ("Çıktılar ölçülebilir hedeflere bağlanmış mı, temenni cümlesi değil mi?", "Bölüm 4"),
        ("Toplumsal/kamusal fayda somut bir kullanıcıya bağlanmış mı?", "Bölüm 4"),
        ("Açık veri / açık kaynak paylaşım taahhüdü var mı?", "Bölüm 4"),
    ])

    kutu(doc, "PUANI NASIL YORUMLARSINIZ", [
        "Her bloğun ortalamasını alın, 4'e bölün, ağırlığıyla çarpın. Örnek: A bloğu ortalaması "
        "3,2 ise → (3,2 / 4) × 35 = 28,0 puan.",
        "Dört bloğun toplamı 100 üzerinden tahmini panel puanınızdır.",
        ("Destek taban puanı her dönem ayrı belirlenir ve ilan edilmez. Deneyim, 70'in altındaki "
         "önerilerin ciddi risk altında olduğunu gösterir — 2 ve altı puan verdiğiniz her satırı "
         "düzeltmeden göndermeyin.", {"bold": True}),
    ], "panelist")

    doc.add_page_break()

    # ---------------------------------------------- 3. Surec akisi
    baslik(doc, "3. BAŞVURU SÜRECİ — sıra ve sorumluluk")
    tablo_yap(doc, ["Adım", "Kim yapar", "Nerede", "Kritik nokta", "Dayanak"], [
        ["1. ARBİS kaydı ve güncelleme", "Yürütücü + ortaklar + danışman",
         "arbis.tubitak.gov.tr", "Bildirimler ARBİS e-postasına gider; adres tebligata elverişli olmalı", "m.4.2.3, m.8.3"],
        ["2. Araştırma önerisinin hazırlanması", "Yürütücü (danışman rehberliğinde)",
         "Resmî .docx formu", "Arial 9, ≤20 sayfa, Türkçe, min karakter sınırları", "m.4.2.1"],
        ["3. Çevrim içi başvuru", "Yürütücü", "tybs.tubitak.gov.tr",
         "Yüklenen dosyanın AÇILDIĞINI ve doğru olduğunu kendiniz kontrol edin", "m.4.3.1"],
        ["4. Son onay", "Yürütücü", "TYBS", "Son gün 17.30'da sistem kapanır", "m.4.3.1"],
        ["5. Danışman onayı", "Akademik danışman", "TYBS (e-posta ile davet)",
         "Danışman iade ederse düzeltip yeniden gönderirsiniz", "m.4.3.4, m.4.3.4.1"],
        ["6. Kuruluş yetkilisi e-imzası", "Rektör / Rektör Yardımcısı", "eimza.tubitak.gov.tr",
         "BU İMZA OLMADAN BAŞVURU GEÇERLİ DEĞİLDİR — takibi SİZİN sorumluluğunuzda", "m.4.3.4.2, m.4.3.4.3"],
        ["7. Ön inceleme", "BİDEB", "—", "Format/belge/karakter kontrolü; burada elenirseniz puan almazsınız", "m.5.1"],
        ["8. Bilimsel değerlendirme", "Panelist / dış danışman", "—", "%35/25/20/20 rubriği", "m.5.2.1"],
        ["9. Etik kurul belgesi (gerekirse)", "Yürütücü", "e-bideb.tubitak.gov.tr",
         "Değerlendirmede istenirse verilen sürede yüklenmezse proje YÜRÜRLÜĞE ALINMAZ", "m.5.3.1"],
        ["10. Sonuçların ilanı", "BİDEB", "e-bideb.tubitak.gov.tr", "Proje süresi bu tarihten itibaren başlar", "m.5.2.8, m.6.1"],
    ], genislikler=[Cm(3.4), Cm(3.0), Cm(3.2), Cm(5.4), Cm(1.8)])

    kutu(doc, "GERİ ÇEKME", [
        "Başvuruyu geri çekmek isterseniz imzalı dilekçeyi bideb2209a@tubitak.gov.tr adresine "
        "gönderirsiniz. Bu işlem çağrı kapanış tarihinden itibaren 1 ay içinde yapılabilir (m.4.3.5).",
    ], "tubitak")

    # ---------------------------------------------- 4. UYZ beyani
    baslik(doc, "4. ÜRETKEN YAPAY ZEKÂ (ÜYZ) KULLANIMI — 2025 çağrısının yeni yükümlülüğü")
    kutu(doc, "ÇAĞRI METNİ MADDE 8.4 ve 8.5", [
        "m.8.4: \"Proje yürütücüsü, proje önerisinin ve projeye özel bütün bilgilerin hazırlanma ve "
        "sunulma sürecinde ve devamında üretken yapay zekâ (ÜYZ) araçlarının kullanımına ilişkin olarak "
        "'TÜBİTAK Destek Süreçlerinde Üretken Yapay Zekânın (ÜYZ) Sorumlu ve Güvenilir Kullanımı "
        "Rehberi'nde yer alan tüm ilke, kural ve yükümlülüklere tam ve eksiksiz biçimde uygun hareket "
        "etmekle yükümlüdür.\"",
        "m.8.5: Akademik danışman, ÜYZ ile üretilen tüm içeriğin doğruluğunu, özgünlüğünü, güncelliğini "
        "ve bilimsel güvenilirliğini kontrol etmek, gerekli düzeltmeleri yapmak ve NİHAİ SORUMLULUĞU "
        "ÜSTLENMEKLE yükümlüdür.",
    ], "tubitak")

    tablo_yap(doc, ["", "Yapmanız gereken"], [
        [KUTU, "Başvurudan önce TÜBİTAK'ın güncel ÜYZ Rehberi'ni tubitak.gov.tr üzerinden okuyun; "
               "beyan yükümlülüğünün o dönemki biçimini teyit edin."],
        [KUTU, "ÜYZ ile üretilmiş her metni satır satır okuyun ve kendi cümlelerinizle yeniden yazın. "
               "Metnin bilimsel içeriğinden SİZ sorumlusunuz."],
        [KUTU, "ÜYZ'nin ürettiği HER kaynağı ayrı ayrı doğrulayın. Var olmayan makale atfı, AYEK "
               "Yönetmeliği kapsamında etik ihlali işlemine yol açar (m.5.2.3, m.8.2)."],
        [KUTU, "ÜYZ'nin ürettiği her sayıyı, oranı ve iddiayı kendi ölçümünüzle veya kaynağıyla "
               "doğrulayın. Doğrulayamadığınızı metinden ÇIKARIN."],
        [KUTU, "Akademik danışmanınıza ÜYZ kullandığınızı açıkça söyleyin — nihai sorumluluk onda."],
        [KUTU, "Kişisel veri, henüz yayımlanmamış özgün fikir veya gizli bilgi içeren metinleri "
               "ÜYZ araçlarına girmeyin."],
    ], genislikler=[Cm(0.8), Cm(16.2)])

    doc.add_page_break()

    # ---------------------------------------------- 5. Sik hatalar
    baslik(doc, "5. PUAN KAYBETTİREN SIK HATALAR")
    tablo_yap(doc, ["Hata", "Neden puan kaybettirir", "Nasıl düzeltilir"], [
        ["\"Bu alanda çok az çalışma vardır\" deyip atıf vermemek",
         "Bilimsel nitelik (%35) doğrudan literatür boşluğunun KANITLANMASINA bakar",
         "Boşluğu 3–5 güncel kaynağa atıfla göster; \"X çalışması Y'yi ele aldı ancak Z'yi ele almadı\" kalıbını kullan"],
        ["Araştırma sorusu yerine yapılacaklar listesi yazmak",
         "Panelist \"bu araştırma değil, geliştirme işi\" der — en sık ret gerekçesi",
         "Soru cümlesi kur: \"... değişkeni ... çıktısını hangi yönde ve ne ölçüde etkiler?\""],
        ["Hipotez yazmamak veya yanlışlanamaz hipotez kurmak",
         "Yöntem (%25) hipotez-test eşleşmesi üzerinden okunur",
         "H0/H1 biçiminde yaz; her hipotezin karşısına uygulanacak istatistiksel testi koy"],
        ["Başarı ölçütünü \"başarılı olması\" gibi niteliksiz yazmak",
         "Form açıkça NİCEL ve izlenebilir ölçüt istiyor (dipnot ***)",
         "Sayı ver: \"etiketleyiciler arası uyum κ ≥ 0,70\", \"makro-F1 ≥ 0,85\""],
        ["Katkı yüzdelerinin toplamının 100 olmaması",
         "Ön incelemede biçim hatası; en azından proje yönetimi puanını kırar",
         "Sütunu toplayın; 100'e tamamlayın"],
        ["Literatür taramasını ayrı iş paketi yapmak",
         "Form bunu AÇIKÇA yasaklıyor (dipnot **)",
         "Literatür taramasını ilgili İP'nin faaliyetleri içine gömün"],
        ["Ürün/marka/kütüphane adı yazmak",
         "m.4.1.8 ihlali — ön incelemede eleme riski",
         "Jenerik karşılık kullanın: \"büyük dil modeli\", \"ilişkisel veritabanı\", \"açık kaynak kütüphane\""],
        ["Yaygın etkide abartılı vaat (\"5 SCI makalesi, 2 patent\")",
         "Panelist gerçekçilik arar; abartı güvenilirliği düşürür",
         "Lisans projesi ölçeğinde yazın: 1–2 ulusal bildiri + 1 makale hedefi gerçekçidir"],
        ["Bütçeye konferans/yayın/patent gideri koymak",
         "m.2.1.2 bunu açıkça yasaklıyor",
         "Bu kalemleri çıkarın; çıktı olarak hedefleyin ama gider yazmayın"],
        ["Riskleri \"proje gecikebilir\" gibi genel yazmak",
         "Proje yönetimi (%20) somut B planı arar",
         "Her riske özgü, uygulanabilir ve ana hedeften saptırmayan bir B planı yazın"],
        ["Ön çalışmayı gizlemek",
         "Form ön çalışma/fizibilite sunulmasını AÇIKÇA bekliyor — gizlemek avantaj kaybıdır",
         "Mevcut prototipi \"ön çalışma bulguları\" başlığıyla sunun, ölçümleriyle birlikte"],
        ["Etik kurul gerekliliğini atlamak",
         "İnsan katılımcılı çalışmada belge gelmezse proje YÜRÜRLÜĞE ALINMAZ (m.5.3.1)",
         "Yöntemde etik kurul onayı/muafiyeti durumunu açıkça belirtin, süreci erken başlatın"],
        ["Son güne bırakmak",
         "Danışman onayı + kuruluş yetkilisi e-imzası ayrı ayrı süre ister",
         "En az 1 hafta önce sistemsel onayı tamamlayın; e-imza takibi sizin sorumluluğunuzda (m.4.3.4.3)"],
    ], genislikler=[Cm(4.6), Cm(5.8), Cm(6.6)])

    doc.add_page_break()

    # ---------------------------------------------- 6. Etik kurul
    baslik(doc, "6. ETİK KURUL — onay mı, muafiyet mi?")
    kutu(doc, "KARAR AĞACI", [
        ("Aşağıdakilerden HERHANGİ BİRİ varsa etik kurul ONAYI gerekir:", {"bold": True}),
        ("Anket, mülakat, odak grup, gözlem gibi yollarla insanlardan veri toplanıyor", {"bullet": True}),
        ("İnsan katılımcılara bir görev/deney uygulanıyor (kullanıcı çalışması dâhil)", {"bullet": True}),
        ("Kişisel veri işleniyor veya kimliği belirlenebilir veri kullanılıyor", {"bullet": True}),
        ("İnsan/hayvan biyolojik materyali kullanılıyor", {"bullet": True}),
        ("Uzman değerlendiriciler bir ölçekle puanlama yapıyor (etiketleyici çalışması)", {"bullet": True}),
        "",
        ("Yalnızca aşağıdakiler varsa genellikle MUAFİYET belgesi yeterlidir:", {"bold": True}),
        ("Tamamen sentetik veya kurgu veri kullanılıyor", {"bullet": True}),
        ("Yalnızca kamuya açık, kişisel veri içermeyen belge/mevzuat metinleri kullanılıyor", {"bullet": True}),
        ("İnsan katılımcı yok, yalnızca ikincil/açık veri üzerinde hesaplama yapılıyor", {"bullet": True}),
        "",
        ("Karar SİZİN üniversitenizin etik kuruluna aittir; yukarıdaki ayrım yol göstericidir. "
         "Kararı erken alın: belge, değerlendirme sonrası verilen kısa sürede yüklenemezse proje "
         "yürürlüğe alınmaz (m.5.3.1).", {"bold": True}),
    ], "panelist")

    # ---------------------------------------------- 7. Sonuc raporu
    baslik(doc, "7. SONUÇ RAPORU İSKELETİ — projeyi kazandıktan sonra")
    kutu(doc, "ÇAĞRI METNİ MADDE 6.7", [
        "\"Projenin bilimsel ve teknik tüm gelişmelerini içeren sonuç raporu ve BİDEB Başvuru ve "
        "İzleme Sistemi'nde yer alan harcama tablosu proje süresi içerisinde proje yürütücüsü "
        "tarafından TÜBİTAK'ça hazırlanmış olan formata uygun olarak yüklenir; akademik danışman "
        "tarafından kontrol edilerek nihai onay verilir.\"",
    ], "tubitak")
    gövde(doc, "TÜBİTAK'ın resmî sonuç raporu formatı şu başlıkları içerir:", bold=True)
    for m in ["Genel bilgiler (projenin konusu, yürütücü, danışman, başlangıç–bitiş tarihleri)",
              "Giriş",
              "Proje konusu hakkında kısa bilgi",
              "Rapor dönemlerinde yapılan çalışmalar",
              "Sonuç",
              "Çıktılar (yayınlar, sunumlar vb.)",
              "Proje ile ilgili harcama kalemleri hakkında ayrıntılı bilgi",
              "Proje yürütücüsü ve danışman imzaları + tarih"]:
        gövde(doc, m, bullet=True)

    tablo_yap(doc, ["", "Sonuç aşaması yükümlülükleri"], [
        [KUTU, "Harcamalar FATURA ile belgelendirildi; faturalar yürütücü veya ortaklar adına kesildi (m.6.2)"],
        [KUTU, "İsimsiz fiş kullanılmadı (akaryakıt fişleri hariç) — geçersiz sayılır ve iadesi istenir (m.6.2)"],
        [KUTU, "Demirbaş alındıysa üniversite envanterine kaydedildi ve kayıt belgesi alındı (m.6.2, m.6.8)"],
        [KUTU, "Kullanılmayan/belgelenemeyen tutar TÜBİTAK hesabına iade edildi (m.6.2)"],
        [KUTU, "Çıktılar (makale, bildiri, kitap bölümü) BİDEB sistemine yüklendi (m.6.6)"],
        [KUTU, "Tüm duyuru ve yayınlarda TÜBİTAK desteği belirtildi ve logo kurallara uygun kullanıldı (m.6.6)"],
        [KUTU, "Sonuç raporu proje süresi içinde sisteme yüklendi (gecikirse gerekçeli dilekçeyle "
               "en fazla 6 ay ek süre istenebilir — m.6.12)"],
    ], genislikler=[Cm(0.8), Cm(16.2)])

    kutu(doc, "PROJEYİ TAMAMLAYINCA AÇILAN KAPILAR (m.7)", [
        "2224-A Yurt Dışı Bilimsel Etkinliklere Katılımı Destekleme Programı",
        "2224-B Yurt İçi Bilimsel Etkinliklere Katılımı Destekleme Programı",
        "2224-D Yurt Dışındaki Bilimsel Eğitim Etkinliklerine Katılımı Destekleme Programı",
    ], "ornek")

    kutu(doc, "İLETİŞİM", [
        "E-posta: bideb2209a@tubitak.gov.tr   ·   Telefon: 444 66 90",
        "Başvuru sistemi: tybs.tubitak.gov.tr   ·   İzleme sistemi: e-bideb.tubitak.gov.tr",
        "ARBİS: arbis.tubitak.gov.tr   ·   E-imza: eimza.tubitak.gov.tr",
    ], "tubitak")

    doc.save(cikti)
    return cikti


if __name__ == "__main__":
    print("YAZILDI:", uret(sys.argv[1]))
