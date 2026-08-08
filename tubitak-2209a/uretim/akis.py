"""Bolum metnini TEK HUCRELIK kutudan cikarip dogrudan govde akisina yazar.

Neden: Formun 1.1 / 1.2 / 2 bolumleri tek satirlik birer tablodur. Ic içerik birkaç
sayfa sürdüğünde Word bu satırı sayfa sınırında beklenmedik biçimde davranıyor ve
yalnızca başlık taşıyan boş sayfalar bırakıyor (ölçüldü: dört sayfa). Ayrıca sayfalarca
süren bir metni çerçeve içinde göstermek tipografik olarak da yanlış: çerçeve bir GİRDİ
ALANI işaretidir, doldurulmuş metnin kendisi değil.

Formun yapısı korunur: başlıklar, sıraları, resmî tablolar (takvim/risk/olanaklar/
yaygın etki) ve kapak aynen kalır. Değişen yalnız serbest metin bölümlerinin kutudan
çıkarılmasıdır.
"""
import re
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from bicim import (bicimle, golge, hucre_kenarlik, hucre_bosluk, satir_turu,
                   LACIVERT, GOVDE, BASLIK_ZEMIN, SERIT_ZEMIN, DOLU_ZEMIN, BOS_ZEMIN)
from docx.shared import RGBColor

BEYAZ = RGBColor(0xFF, 0xFF, 0xFF)


def _tasi(hedef_el, yeni_el):
    """Yeni ogeyi hedef ogenin ONUNE tasir."""
    hedef_el.addprevious(yeni_el)


def _boru_tablosu_akis(doc, hedef_el, satirlar, toplam_genislik):
    sutun = max(len(r) for r in satirlar)
    satirlar = [r + [""] * (sutun - len(r)) for r in satirlar]

    t = doc.add_table(rows=len(satirlar), cols=sutun)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False

    uzunluk = [max(len(r[c]) for r in satirlar) for c in range(sutun)]
    toplam = sum(uzunluk) or 1
    oranlar = [min(max(u / toplam, 0.06), 0.40) for u in uzunluk]
    olcek = sum(oranlar)
    enler = [Cm(toplam_genislik.cm * (o / olcek)) for o in oranlar]

    for ri, satir in enumerate(satirlar):
        for ci, deger in enumerate(satir):
            c = t.rows[ri].cells[ci]
            c.width = enler[ci]
            hucre_kenarlik(c)
            hucre_bosluk(c, ust=30, alt=30, sol=70, sag=70)
            c._element.clear_content()
            p = c.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
            if ri == 0:
                golge(c, BASLIK_ZEMIN)
                bicimle(p.add_run(deger), size=7.5, bold=True, renk=BEYAZ)
            else:
                if ri % 2 == 0:
                    golge(c, SERIT_ZEMIN)
                bicimle(p.add_run(deger), size=7.5, bold=(ci == 0),
                        renk=LACIVERT if ci == 0 else GOVDE)
    _tasi(hedef_el, t._tbl)

    bosluk = doc.add_paragraph()
    bosluk.paragraph_format.space_after = Pt(5)
    bicimle(bosluk.add_run(""), size=5)
    _tasi(hedef_el, bosluk._p)
    return t


def zengin_akis(doc, hedef_el, metin, taban=9, genislik=Cm(16.4)):
    """Metni hedef ogenin onune, bicimli paragraflar ve gercek tablolar olarak yazar."""
    satirlar = metin.split("\n")
    i = 0
    while i < len(satirlar):
        tur = satir_turu(satirlar[i])
        if tur == "bos":
            i += 1
            continue

        if tur == "boru":
            blok = []
            while i < len(satirlar) and satir_turu(satirlar[i]) == "boru":
                blok.append([h.strip() for h in satirlar[i].split("|")])
                i += 1
            _boru_tablosu_akis(doc, hedef_el, blok, genislik)
            continue

        t = satirlar[i].strip()
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.widow_control = True

        if tur == "baslik":
            pf.space_before = Pt(7)
            pf.space_after = Pt(2)
            pf.keep_with_next = True
            bicimle(p.add_run(t), size=taban, bold=True, renk=LACIVERT)
        elif tur == "madde":
            govde = re.sub(r"^[-•]\s*", "", t)
            pf.space_after = Pt(1)
            pf.left_indent = Pt(14)
            pf.first_line_indent = Pt(-9)
            if len(govde) > 120:
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            bicimle(p.add_run("•  "), size=taban)
            bicimle(p.add_run(govde), size=taban)
        elif tur == "sirali":
            m = re.match(r"^(\(\w{1,3}\))\s*(.*)$", t, re.S)
            pf.space_after = Pt(1)
            pf.left_indent = Pt(18)
            pf.first_line_indent = Pt(-18)
            if len(t) > 120:
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            bicimle(p.add_run(m.group(1) + " "), size=taban, bold=True)
            bicimle(p.add_run(m.group(2)), size=taban)
        else:
            pf.space_after = Pt(2)
            if len(t) > 120:
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            bicimle(p.add_run(t), size=taban)

        _tasi(hedef_el, p._p)
        i += 1


def kutuyu_akisa_cevir(doc, tablo, metin, taban=9):
    """Tek hucrelik icerik kutusunu kaldirir, metni onun yerine govde akisi olarak yazar."""
    hedef = tablo._tbl
    zengin_akis(doc, hedef, metin, taban=taban)
    hedef.getparent().remove(hedef)


# ------------------------------------------------------------ ŞEKİLLER (akış sürümü)
def _yazi(doc, hedef_el, metin, size=8, bold=False, italic=False, renk=None,
          once=8, sonra=3, birlikte=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(once)
    pf.space_after = Pt(sonra)
    pf.keep_with_next = birlikte
    bicimle(p.add_run(metin), size=size, bold=bold, italic=italic, renk=renk or GOVDE)
    _tasi(hedef_el, p._p)
    return p


def _hucre(t, ri, ci, metin, size=7.5, bold=False, renk=None, orta=False, alt_metin=None):
    c = t.rows[ri].cells[ci]
    c._element.clear_content()
    p = c.add_paragraph()
    p.paragraph_format.space_after = Pt(0 if alt_metin else 1)
    p.paragraph_format.space_before = Pt(1)
    if orta:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    bicimle(p.add_run(str(metin)), size=size, bold=bold, renk=renk or GOVDE)
    if alt_metin:
        p2 = c.add_paragraph()
        p2.paragraph_format.space_after = Pt(1)
        bicimle(p2.add_run(alt_metin), size=6.5, italic=True,
                renk=RGBColor(0x55, 0x55, 0x55))
    return c


def faktoriyel_akis(doc, hedef_el):
    """Şekil 1 — D1'in 4 (mimari) × 3 (bozulma) gövdesi, gövde akışında."""
    _yazi(doc, hedef_el, "Şekil 1. D1 deneyinin faktöriyel gövdesi "
                         "(4 mimari kolu × 3 girdi bozulma düzeyi, belge-içi tekrarlı ölçüm)",
          size=8.5, bold=True, renk=LACIVERT, once=10, sonra=4, birlikte=True)

    kollar = [("(a) Kural katmanı", "deterministik; model çağrısı yok"),
              ("(b) Kısıtsız üretim", "alt sınır referansı"),
              ("(b2) Literatür tabanı", "şemaya bağlı + öz-eleştiri turu"),
              ("(c) Kısıtlı hibrit", "kapalı küme + kanıt alıntısı")]
    duzeyler = ["Temiz (KHO ≈ 0)", "Orta (KHO ≈ 0,05)", "Yüksek (KHO ≈ 0,12)"]

    t = doc.add_table(rows=5, cols=4)
    t.autofit = False
    enler = [Cm(5.6), Cm(3.6), Cm(3.6), Cm(3.6)]

    _hucre(t, 0, 0, "Mimari kolu ↓   |   Bozulma düzeyi →", bold=True, renk=BEYAZ)
    golge(t.rows[0].cells[0], BASLIK_ZEMIN)
    for j, d in enumerate(duzeyler):
        _hucre(t, 0, j + 1, d, bold=True, renk=BEYAZ, orta=True)
        golge(t.rows[0].cells[j + 1], BASLIK_ZEMIN)

    for i, (ad, aciklama) in enumerate(kollar, start=1):
        _hucre(t, i, 0, ad, bold=True, renk=LACIVERT, alt_metin=aciklama)
        for j in range(1, 4):
            _hucre(t, i, j, "240 belge", orta=True)
        if i % 2 == 0:
            for j in range(4):
                golge(t.rows[i].cells[j], SERIT_ZEMIN)

    for r in t.rows:
        for j, c in enumerate(r.cells):
            c.width = enler[j]
            hucre_kenarlik(c)
            hucre_bosluk(c, ust=35, alt=35)
    _tasi(hedef_el, t._tbl)

    _yazi(doc, hedef_el, "Gövde 240 × 4 × 3 = 2.880 gözlem; sağlayıcı bloğu 80 belgelik alt "
                         "kümede iki ek sağlayıcıda tekrarlanır (+1.440), toplam 4.320. "
                         "(a) kolu model çağrısı içermez: 720 gözlem deterministiktir. "
                         "KHO: karakter hata oranı.",
          size=7, italic=True, renk=RGBColor(0x44, 0x44, 0x44), once=3, sonra=8)


def gantt_akis(doc, hedef_el, takvim):
    """Şekil 2 — 12 aylık iş paketi zaman çizelgesi, gövde akışında."""
    _yazi(doc, hedef_el, "Şekil 2. On iki aylık iş paketi zaman çizelgesi ve katkı ağırlıkları",
          size=8.5, bold=True, renk=LACIVERT, once=10, sonra=4, birlikte=True)

    t = doc.add_table(rows=len(takvim) + 1, cols=14)
    t.autofit = False
    basliklar = ["İş paketi"] + [str(a) for a in range(1, 13)] + ["Katkı"]
    for j, b in enumerate(basliklar):
        _hucre(t, 0, j, b, size=7, bold=True, renk=BEYAZ, orta=True)
        golge(t.rows[0].cells[j], BASLIK_ZEMIN)

    for i, ip in enumerate(takvim, start=1):
        ad = str(ip.get("ad", ""))
        kisa = ad.split("—")[0].strip() if "—" in ad else ad[:6]
        baslik = ad.split("—", 1)[1].strip() if "—" in ad else ad
        _hucre(t, i, 0, kisa, size=7.5, bold=True, renk=LACIVERT,
               alt_metin=baslik[:70] + ("…" if len(baslik) > 70 else ""))
        m = re.search(r"(\d+)\s*-\s*(\d+)", str(ip.get("tarih_araligi", "")))
        bas, son = (int(m.group(1)), int(m.group(2))) if m else (1, 1)
        for ay in range(1, 13):
            _hucre(t, i, ay, "", size=7, orta=True)
            golge(t.rows[i].cells[ay], DOLU_ZEMIN if bas <= ay <= son else BOS_ZEMIN)
        _hucre(t, i, 13, f"%{ip.get('katki')}", size=7.5, bold=True,
               renk=LACIVERT, orta=True)

    for r in t.rows:
        for j, c in enumerate(r.cells):
            c.width = Cm(4.5) if j == 0 else (Cm(1.3) if j == 13 else Cm(0.72))
            hucre_kenarlik(c)
            hucre_bosluk(c, ust=22, alt=22, sol=30, sag=30)

    # Sekil sayfa siniri asabiliyor; asarsa ay basligi devam sayfasinda da gorunsun,
    # yoksa okuyucu hangi hucrenin hangi ay oldugunu bilemez.
    trPr = t.rows[0]._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader")
    e.set(qn("w:val"), "true")
    trPr.append(e)

    _tasi(hedef_el, t._tbl)

    toplam = sum(x.get("katki") or 0 for x in takvim)
    # "3.1" numarasi KALDIRILDI: formda o numarayla bir baslik yok, atif havada kaliyordu.
    _yazi(doc, hedef_el, f"Koyu hücreler iş paketinin etkin olduğu ayları gösterir. Katkı "
                         f"sütununun toplamı %{toplam}'dür. Şekil, Çalışma Takvimi "
                         f"tablosunun görsel özetidir; ayrıntılı faaliyet, sorumlu ve başarı "
                         f"ölçütü bilgisi o tablodadır.",
          size=7, italic=True, renk=RGBColor(0x44, 0x44, 0x44), once=3, sonra=6)
