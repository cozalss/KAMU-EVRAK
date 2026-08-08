"""Iki yardimci belge uretir:

  B. 2209A_REHBERLI_SABLON.docx  — her bolumde 4 katman:
       [TUBITAK] resmi aciklama · [PANELIST NE ARAR] rubrik ipucu ·
       [ORNEK] doldurulmus metin · [SIZ YAZIN] bos alan + karakter hedefi
  C. 2209A_KONTROL_LISTESI.docx  — on-inceleme eleme listesi, rubrik oz-puanlama,
       UYZ beyan rehberi, surec akisi, sik hatalar, sonuc raporu iskeleti

Tum icerik cagri duyurusu (2025) ve resmi arastirma onerisi formundan tureltilmistir.
"""
import sys, os, json
sys.stdout.reconfigure(encoding="utf-8")

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ARIAL = "Arial"
W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

RENK = {
    "tubitak": RGBColor(0x1F, 0x49, 0x7D),   # koyu mavi  (resmi aciklama)
    "panelist": RGBColor(0x8B, 0x25, 0x00),  # koyu kiremit (rubrik ipucu)
    "ornek": RGBColor(0x1E, 0x5E, 0x20),   # koyu yesil (ornek metin)
    "siz": RGBColor(0x40, 0x40, 0x40),   # gri (bos alan)
    "uyari": RGBColor(0xA0, 0x00, 0x00),
    "govde": RGBColor(0x11, 0x11, 0x11),
}
ZEMIN = {"tubitak": "EAF0F7", "panelist": "FBEDE8", "ornek": "EAF3EB",
         "siz": "F5F5F5", "uyari": "FDE8E8"}


def _font(run, size=9, bold=False, italic=False, renk=None, ad=ARIAL):
    run.font.name = ad
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if renk is not None:
        run.font.color.rgb = renk
    rpr = run._element.get_or_add_rPr()
    rf = rpr.get_or_add_rFonts()
    for a in ("ascii", "hAnsi", "cs", "eastAsia"):
        rf.set(W_NS + a, ad)
    return run


def _golge(hucre, hex_renk):
    tcPr = hucre._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_renk)
    tcPr.append(shd)


def _kenar(hucre, renk="BFBFBF", kalinlik="6"):
    tcPr = hucre._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for yon in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{yon}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), kalinlik)
        e.set(qn("w:color"), renk)
        borders.append(e)
    tcPr.append(borders)


def kutu(doc, etiket, satirlar, tur, kutu_genisligi=Cm(17.0)):
    """Renkli, tek hucreli bilgi kutusu ekler."""
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    c = t.rows[0].cells[0]
    c.width = kutu_genisligi
    _golge(c, ZEMIN[tur])
    _kenar(c)
    c._element.clear_content()

    p = c.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    _font(p.add_run(etiket), size=8, bold=True, renk=RENK[tur])

    for s in satirlar:
        if isinstance(s, tuple):
            metin, opt = s
        else:
            metin, opt = s, {}
        p = c.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        if len(str(metin)) > 110:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if opt.get("bullet"):
            metin = "•  " + str(metin)
            p.paragraph_format.left_indent = Pt(10)
        _font(p.add_run(str(metin)), size=opt.get("size", 9),
              bold=opt.get("bold", False), italic=opt.get("italic", False),
              renk=RENK["govde"] if tur == "ornek" else RENK[tur])
    doc.add_paragraph()
    return t


def bos_alan(doc, hedef_karakter, ipucu="", etiket=None, satir=6):
    t = doc.add_table(rows=1, cols=1)
    c = t.rows[0].cells[0]
    c.width = Cm(17.0)
    _golge(c, ZEMIN["siz"])
    _kenar(c, renk="9E9E9E")
    c._element.clear_content()
    p = c.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    if etiket is None:
        etiket = (f"▼ BURAYA SİZ YAZIN — hedef en az {hedef_karakter} karakter "
                  f"(TÜBİTAK alt sınırı 350; altında kalırsanız ön incelemede ELENİRSİNİZ)")
    _font(p.add_run(etiket), size=8, bold=True, renk=RENK["siz"])
    if ipucu:
        p = c.add_paragraph()
        _font(p.add_run(ipucu), size=8, italic=True, renk=RENK["siz"])
    for _ in range(satir):
        p = c.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        _font(p.add_run(" "), size=9)
    doc.add_paragraph()


def baslik(doc, metin, seviye=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10 if seviye == 1 else 6)
    p.paragraph_format.space_after = Pt(4)
    _font(p.add_run(metin), size=13 if seviye == 1 else 11, bold=True,
          renk=RGBColor(0x0D, 0x2B, 0x45))
    if seviye == 1:
        pPr = p._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "8")
        bottom.set(qn("w:color"), "0D2B45")
        bdr.append(bottom)
        pPr.append(bdr)
    return p


def gövde(doc, metin, size=9, bold=False, bullet=False, renk=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    if bullet:
        metin = "•  " + metin
        p.paragraph_format.left_indent = Pt(12)
    if len(metin) > 110:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _font(p.add_run(metin), size=size, bold=bold, renk=renk or RENK["govde"])
    return p


def sayfa_kur(doc):
    for s in doc.sections:
        s.top_margin = Cm(1.9)
        s.bottom_margin = Cm(1.6)
        s.left_margin = Cm(2.0)
        s.right_margin = Cm(2.0)
    st = doc.styles["Normal"]
    st.font.name = ARIAL
    st.font.size = Pt(9)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), ARIAL)


def tablo_yap(doc, basliklar, satirlar, genislikler=None, baslik_zemin="0D2B45"):
    t = doc.add_table(rows=1, cols=len(basliklar))
    t.style = "Table Grid"
    if genislikler:
        t.autofit = False
        for i, g in enumerate(genislikler):
            for row in t.rows:
                row.cells[i].width = g
    for c, b in zip(t.rows[0].cells, basliklar):
        _golge(c, baslik_zemin)
        c._element.clear_content()
        p = c.add_paragraph()
        _font(p.add_run(str(b)), size=8, bold=True, renk=RGBColor(0xFF, 0xFF, 0xFF))
    for s in satirlar:
        row = t.add_row()
        if genislikler:
            for i, g in enumerate(genislikler):
                row.cells[i].width = g
        for c, v in zip(row.cells, s):
            c._element.clear_content()
            p = c.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            _font(p.add_run(str(v)), size=8)
    doc.add_paragraph()
    return t
